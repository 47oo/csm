"""F037 解析异常收敛为 422 ``IMPORT_INVALID_FILE`` — 独立 TEST 执行。

覆盖 `docs/features/F037/test-report.md` §1 `TD-F037-1` 的 U / I / S 用例，并补齐
`docs/features/F037/architecture.md` §15.8（输入 ③ / 白名单 11 类 / `SyntaxError`
负向守卫 / 静态守卫）与 §15.10.7（`U13` / `U4` 11 类 / `U14` / `I22` / `S7`）。

分层（architecture §10.1）：

* **U** — 直接调用 `app.services.resource_import` 的纯函数，不需要数据库；
* **I** — `auth_app`（真实 `create_app`）+ 真实 PostgreSQL（`TEST_DATABASE_URL`，
  DB-D10 双守卫）；
* **S** — 对生产源码做静态/AST 守卫（`git diff` 类守卫在 test-report §2 以命令记录）。

全部输入在内存中构造（`io.BytesIO` / `zipfile`），**不读写任何磁盘上的 xlsx 或
被跟踪的 artifacts 文件**，可独立重复执行。

本模块不导入 `app.core.access_scope` 符号，故按 F030 §9.1 不需要
`pytest.mark.real_scope`（由 `test_f030_test_harness.py` 的 AST 门禁复核）。
"""

from __future__ import annotations

import ast
import csv
import io
import logging
import os
import re
import xml.etree.ElementTree as ElementTree
import zipfile
import zlib
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import text

from app.api.v1 import imports as imports_api
from app.core.config import Settings
from app.core.errors import (
    ImportFileTooLarge,
    ImportInvalidFile,
    ImportRowLimitExceeded,
    ImportValidationFailed,
)
from app.services import resource_import

from integration.f015_helpers import role_client_factory  # noqa: F401  (reused fixture)

# --- constants ---------------------------------------------------------------

RES_VALIDATE = "/api/v1/imports/resources/validate"
RES_WRITE = "/api/v1/imports/resources"
UNI_VALIDATE = "/api/v1/imports/resource-interfaces-ips/validate"
UNI_WRITE = "/api/v1/imports/resource-interfaces-ips"
ALL_PATHS = (RES_VALIDATE, RES_WRITE, UNI_VALIDATE, UNI_WRITE)
WRITE_PATHS = (RES_WRITE, UNI_WRITE)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# F014 BM 14 列 / F023 v2 26 列冻结列序（Contract §6；零变更由 S5 复核）。
BM_HEADERS = [
    "cluster",
    "hostname",
    "resource_type",
    "role",
    "status",
    "manufacturer",
    "model",
    "serial_number",
    "cpu_model",
    "cpu_sockets",
    "cpu_cores",
    "memory_mb",
    "gpu_model",
    "gpu_count",
]
W26_HEADERS = [
    "cluster",
    "hostname",
    "resource_type",
    "role",
    "status",
    "host",
    "vcpu",
    "memory_mb",
    "disk_gb",
    "hypervisor",
    "manufacturer",
    "model",
    "serial_number",
    "cpu_model",
    "cpu_sockets",
    "cpu_cores",
    "gpu_model",
    "gpu_count",
    "name",
    "interface_type",
    "mac_address",
    "speed_mbps",
    "mtu",
    "enabled",
    "network",
    "ip",
]

# 载体形状常量（与 F034 `importCarrier.ts` / test_f034_csv_text_carrier.py 同值）。
PASTED_FILE_NAME = "csm-pasted-import.csv"
PASTED_MIME = "text/csv;charset=utf-8"
FILE_NAME = "f037-file-carrier.csv"
FILE_MIME = "text/csv"

MARKER = "F037-MARKER-5c1d9b7a2e4f4ba7"

COUNT_TABLES = (
    "clusters",
    "resources",
    "bare_metal_details",
    "virtual_machine_details",
    "interfaces",
    "ip_addresses",
    "audit_logs",
)

# F014 §1.2 / F023 v2 §1.2 已签核 reason 集合（本 Feature 不得新增取值）。
SIGNED_REASONS = {
    "unreadable",
    "unsupported_format",
    "header_missing",
    "missing_columns",
    "unknown_columns",
    "column_order",
    "empty_file",
}

UNREADABLE_BODY = {
    "error": {
        "code": "IMPORT_INVALID_FILE",
        "message": "Invalid import file",
        "details": {"reason": "unreadable"},
    }
}

# `_PARSE_FAILURES` 为 8 项（零 diff）；`.xlsx` 局部再并 `InvalidFileException`、
# 条件 `lxml.etree.LxmlError`、`OSError`（DR-F037-08）⇒ 有效 11 类。
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"not-a-real-xls"


# --- payload builders (all in memory) ----------------------------------------


def headers_for(path: str) -> list[str]:
    """每个端点使用自己的冻结模板表头（BM 14 / 26 列）。"""
    return W26_HEADERS if "resource-interfaces-ips" in path else BM_HEADERS


def oversized_field_csv(
    headers: list[str], *, pad: int = 131_073, marker: str = ""
) -> bytes:
    """输入 ①：单字段超过 ``csv.field_size_limit()``（本机 131072）。"""
    row = [""] * len(headers)
    row[0] = "F37A"
    row[-1] = marker + "x" * max(0, pad - len(marker))
    return (",".join(headers) + "\n" + ",".join(row) + "\n").encode("utf-8")


def header_only_csv(headers: list[str]) -> bytes:
    return (",".join(headers) + "\n").encode("utf-8")


def _rows(headers: list[str], rows: list[dict]) -> list[list[str]]:
    return [
        ["" if rec.get(name) is None else str(rec.get(name)) for name in headers]
        for rec in rows
    ]


def csv_bytes(headers: list[str], rows: list[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(headers)
    for row in _rows(headers, rows):
        writer.writerow(row)
    return buffer.getvalue().encode("utf-8")


def xlsx_bytes(headers: list[str], rows: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in _rows(headers, rows):
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sample_rows(headers: list[str]) -> list[dict]:
    row = {"cluster": "F37A", "hostname": "f37-x", "resource_type": "BARE_METAL"}
    return [row]


def _rewrite_zip(raw: bytes, member: str, old: bytes | None, new: bytes) -> bytes:
    """重写 zip 内某个成员（``old is None`` 表示删除该成员）。"""
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            if item.filename == member:
                if old is None:
                    continue
                payload = src.read(item.filename)
                assert old in payload, f"{member}: pattern {old!r} missing"
                dst.writestr(item, payload.replace(old, new))
            else:
                dst.writestr(item, src.read(item.filename))
    return out.getvalue()


def _corrupt_deflate(raw: bytes, member: str) -> bytes:
    """把某成员的 deflate 流首字节翻转（zlib.error 路径，真实 payload）。"""
    info = zipfile.ZipFile(io.BytesIO(raw)).getinfo(member)
    data = bytearray(raw)
    name_len = int.from_bytes(data[info.header_offset + 26 : info.header_offset + 28], "little")
    extra_len = int.from_bytes(data[info.header_offset + 28 : info.header_offset + 30], "little")
    start = info.header_offset + 30 + name_len + extra_len
    data[start] ^= 0xFF
    return bytes(data)


def corrupt_sheet_tail(headers: list[str], *, marker: str = "") -> bytes:
    """输入 ②：`xl/worksheets/sheet1.xml` 结束标签损坏（延迟到 iter_rows 才失败）。"""
    rows = _sample_rows(headers)
    if marker:
        rows[0]["role"] = marker
    good = xlsx_bytes(headers, rows)
    return _rewrite_zip(
        good, "xl/worksheets/sheet1.xml", b"</worksheet>", b"</workshet>"
    )


def corrupt_manifest(headers: list[str]) -> bytes:
    """输入 ③：`[Content_Types].xml` 结束标签损坏（CONTRACT TD-I-1 路径）。"""
    good = xlsx_bytes(headers, _sample_rows(headers))
    return _rewrite_zip(good, "[Content_Types].xml", b"</Types>", b"</Typse>")


def manifest_without_workbook_part() -> bytes:
    """BE-01：`[Content_Types].xml` 可读但未声明 workbook 部件（OSError 路径）。"""
    content_types = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        b'<Default Extension="rels" ContentType="application/vnd.openxmlformats-'
        b'package.relationships+xml"/>'
        b'<Default Extension="xml" ContentType="application/xml"/>'
        b"</Types>"
    )
    relationships = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
        b'relationships"/>'
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        dst.writestr("[Content_Types].xml", content_types)
        dst.writestr("_rels/.rels", relationships)
    return out.getvalue()


def zero_member_zip() -> bytes:
    """零成员 zip（magic 为 `PK\\x05\\x06`，落在 PK 闸门 ⇒ unsupported_format）。"""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w"):
        pass
    return out.getvalue()


def xlsx_without_sheet_part(headers: list[str]) -> bytes:
    """合法容器但无 worksheet 部件 ⇒ `workbook.worksheets` 为空 ⇒ empty_file。"""
    good = xlsx_bytes(headers, _sample_rows(headers))
    return _rewrite_zip(good, "xl/worksheets/sheet1.xml", None, b"")


def xlsx_invalid_coordinate(headers: list[str]) -> bytes:
    """非法单元格坐标 ⇒ openpyxl 抛 ValueError（真实 payload）。"""
    good = xlsx_bytes(headers, _sample_rows(headers))
    return _rewrite_zip(good, "xl/worksheets/sheet1.xml", b'<c r="A2"', b'<c r="??"')


def deflate_corrupted_xlsx(headers: list[str]) -> bytes:
    """sheet 部件 deflate 流损坏 ⇒ zlib.error（真实 payload）。"""
    good = xlsx_bytes(headers, _sample_rows(headers))
    return _corrupt_deflate(good, "xl/worksheets/sheet1.xml")


def expected_manifest_error_type() -> type[BaseException]:
    """输入 ③ 的裸异常类型随 `openpyxl.LXML` 分支（§15.8.1，不硬编码后端）。"""
    import openpyxl

    if openpyxl.LXML:  # pragma: no cover - branch depends on the environment
        from lxml.etree import XMLSyntaxError

        return XMLSyntaxError
    return ElementTree.ParseError


def lxml_available() -> bool:
    try:
        import lxml.etree  # noqa: F401
    except ImportError:  # pragma: no cover - mirrors the implementation guard
        return False
    return True


# --- HTTP / assertion helpers ------------------------------------------------


def post(client, path: str, data: bytes, fmt: str = "csv", carrier: str = "file"):
    if fmt == "csv":
        name, mime = (
            (PASTED_FILE_NAME, PASTED_MIME) if carrier == "text" else (FILE_NAME, FILE_MIME)
        )
    else:
        name, mime = ("f037-file-carrier.xlsx", XLSX_MIME)
    return client.post(
        path, files={"file": (name, data, mime)}, data={"format": fmt}
    )


def assert_unreadable(resp, label: str = "") -> None:
    """EB：8 例共用断言包（A1–A8，无 500 / 413 / 2xx）。"""
    assert resp.status_code == 422, f"{label}: {resp.status_code} {resp.text[:400]}"
    assert resp.status_code not in (200, 201), f"{label}: unexpected success"
    body = resp.json()
    assert set(body) == {"error"}, f"{label}: {body}"
    assert set(body["error"]) == {"code", "message", "details"}, f"{label}: {body}"
    assert body == UNREADABLE_BODY, f"{label}: {body}"


def assert_invalid_file(resp, reason: str, label: str = "") -> None:
    assert resp.status_code == 422, f"{label}: {resp.status_code} {resp.text[:400]}"
    error = resp.json()["error"]
    assert error["code"] == "IMPORT_INVALID_FILE", f"{label}: {error}"
    assert error["message"] == "Invalid import file", f"{label}: {error}"
    assert error["details"].get("reason") == reason, f"{label}: {error}"
    assert error["details"].get("reason") != "unreadable" or reason == "unreadable"


def counts(session) -> dict[str, int]:
    return {
        table: int(session.execute(text(f"SELECT count(*) FROM {table}")).scalar_one())
        for table in COUNT_TABLES
    }


def _raising(exc: BaseException):
    def _factory(*_args, **_kwargs):
        raise exc

    return _factory


# =============================================================================
# U 层 — 输入有效性与白名单语义
# =============================================================================


@pytest.mark.parametrize(
    "headers, expected_columns",
    [(BM_HEADERS, 14), (W26_HEADERS, 26)],
    ids=["bm14", "w26"],
)
def test_td_f037_u1_oversized_csv_really_fails_during_parse(headers, expected_columns):
    """TD-F037-U1：输入 ① 有效性硬前置（构造失效即 BLOCKED，不得改判 PASS）。"""
    raw = oversized_field_csv(headers)
    assert len(headers) == expected_columns
    assert 131_072 < len(raw) <= 5_242_880, len(raw)
    with pytest.raises(csv.Error) as excinfo:
        list(csv.reader(io.StringIO(raw.decode("utf-8-sig"), newline="")))
    assert "field larger than field limit" in str(excinfo.value)

    with pytest.raises(ImportInvalidFile) as perr:
        resource_import.parse_records(raw, "csv")
    assert perr.value.details == {"reason": "unreadable"}
    assert isinstance(perr.value.__cause__, csv.Error)


def test_td_f037_u2_corrupt_sheet_tail_really_fails_during_parse():
    """TD-F037-U2：输入 ② 有效性硬前置（load_workbook 成功、iter_rows 抛错）。"""
    raw = corrupt_sheet_tail(BM_HEADERS)
    assert raw.startswith(b"PK\x03\x04"), "PK 魔数闸门不得抢先拦截"
    assert 0 < len(raw) < 5_242_880, len(raw)

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Sheet"]
        with pytest.raises(ElementTree.ParseError):
            list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()

    with pytest.raises(ImportInvalidFile) as perr:
        resource_import.parse_records(raw, "xlsx")
    assert perr.value.details == {"reason": "unreadable"}
    assert isinstance(perr.value.__cause__, ElementTree.ParseError)


def test_td_f037_u15_manifest_corruption_converges_on_both_backends():
    """architecture §15.8.1（输入 ③）：裸类型随 openpyxl.LXML，收敛结果不变。"""
    import openpyxl

    raw = corrupt_manifest(BM_HEADERS)
    assert raw.startswith(b"PK\x03\x04")
    expected = expected_manifest_error_type()

    with pytest.raises(expected):
        load_workbook(io.BytesIO(raw), read_only=True, data_only=True)

    with pytest.raises(ImportInvalidFile) as perr:
        resource_import.parse_records(raw, "xlsx")
    assert perr.value.details == {"reason": "unreadable"}
    assert isinstance(perr.value.__cause__, expected), (
        f"openpyxl.LXML={openpyxl.LXML} expected {expected!r}, "
        f"got {type(perr.value.__cause__)!r}"
    )


def test_td_f037_u3_whitelisted_real_payloads_keep_reason_and_cause():
    """TD-F037-U3：真实 payload 全部收敛为 unreadable 且保留原因链。"""
    real_cases = [
        ("csv-oversized", oversized_field_csv(BM_HEADERS), "csv", csv.Error),
        ("sheet-xml", corrupt_sheet_tail(BM_HEADERS), "xlsx", ElementTree.ParseError),
        ("manifest-xml", corrupt_manifest(BM_HEADERS), "xlsx", expected_manifest_error_type()),
        ("deflate", deflate_corrupted_xlsx(BM_HEADERS), "xlsx", zlib.error),
        (
            "not-a-zip",
            b"PK\x03\x04not-a-zip",
            "xlsx",
            zipfile.BadZipFile,
        ),
        (
            "missing-content-types",
            _rewrite_zip(
                xlsx_bytes(BM_HEADERS, _sample_rows(BM_HEADERS)),
                "[Content_Types].xml",
                None,
                b"",
            ),
            "xlsx",
            KeyError,
        ),
        ("no-workbook-part", manifest_without_workbook_part(), "xlsx", OSError),
        (
            "invalid-coordinate",
            xlsx_invalid_coordinate(BM_HEADERS),
            "xlsx",
            ValueError,
        ),
    ]
    for label, raw, fmt, cause_type in real_cases:
        with pytest.raises(ImportInvalidFile) as perr:
            resource_import.parse_records(raw, fmt)
        assert perr.value.details == {"reason": "unreadable"}, label
        assert perr.value.status_code == 422, label
        assert type(perr.value.__cause__) is cause_type, (
            f"{label}: expected {cause_type!r}, got {type(perr.value.__cause__)!r}"
        )


def test_td_f037_u4_whitelist_covers_eleven_classes(monkeypatch):
    """TD-F037-U4（§15.10.7-2）：11 类白名单参数化，真实 payload 优先、注入兜底。"""
    good = xlsx_bytes(BM_HEADERS, _sample_rows(BM_HEADERS))
    real = {
        "csv.Error": (lambda: oversized_field_csv(BM_HEADERS), "csv", csv.Error),
        "BadZipFile": (lambda: b"PK\x03\x04not-a-zip", "xlsx", zipfile.BadZipFile),
        "zlib.error": (
            lambda: deflate_corrupted_xlsx(BM_HEADERS),
            "xlsx",
            zlib.error,
        ),
        "ElementTree.ParseError": (
            lambda: corrupt_sheet_tail(BM_HEADERS),
            "xlsx",
            ElementTree.ParseError,
        ),
        "KeyError": (
            lambda: _rewrite_zip(good, "[Content_Types].xml", None, b""),
            "xlsx",
            KeyError,
        ),
        "ValueError": (
            lambda: xlsx_invalid_coordinate(BM_HEADERS),
            "xlsx",
            ValueError,
        ),
        "OSError": (manifest_without_workbook_part, "xlsx", OSError),
    }
    injected = [
        zipfile.LargeZipFile("zip64 overflow"),
        NotImplementedError("unsupported compression method"),
        InvalidFileException("invalid container"),
    ]
    if lxml_available():
        from lxml.etree import LxmlError, XMLSyntaxError

        real["LxmlError"] = (lambda: corrupt_manifest(BM_HEADERS), "xlsx", XMLSyntaxError)
        injected.append(LxmlError("injected lxml error"))

    covered: list[str] = []
    for name, (builder, fmt, cause_type) in real.items():
        with pytest.raises(ImportInvalidFile) as perr:
            resource_import.parse_records(builder(), fmt)
        assert perr.value.details == {"reason": "unreadable"}, name
        assert type(perr.value.__cause__) is cause_type, name
        covered.append(name)

    for exc in injected:
        with monkeypatch.context() as patch:
            patch.setattr("openpyxl.load_workbook", _raising(exc))
            with pytest.raises(ImportInvalidFile) as perr:
                resource_import.parse_records(good, "xlsx")
        assert perr.value.details == {"reason": "unreadable"}, type(exc).__name__
        assert type(perr.value.__cause__) is type(exc)
        covered.append(f"injected:{type(exc).__name__}")

    assert "csv.Error" in covered and "OSError" in covered
    assert len(covered) >= 10, covered  # 11 类（无 lxml 环境少 LxmlError）


def test_td_f037_u5_app_error_reason_fidelity(monkeypatch):
    """TD-F037-U5（U-c）：既有 AppError 原样透传，绝不改写为 unreadable。"""
    with pytest.raises(ImportInvalidFile) as empty_csv:
        resource_import.parse_records(b"", "csv")
    assert empty_csv.value.details == {"reason": "empty_file"}

    with pytest.raises(ImportInvalidFile) as empty_xlsx:
        resource_import.parse_records(b"", "xlsx")
    assert empty_xlsx.value.details == {"reason": "empty_file"}

    with pytest.raises(ImportInvalidFile) as ole2:
        resource_import.parse_records(OLE2_MAGIC, "xlsx")
    assert ole2.value.details == {"reason": "unsupported_format"}

    app_errors = [
        ImportInvalidFile("column_order"),
        ImportInvalidFile("missing_columns", missing_columns=["gpu_count"]),
        ImportInvalidFile("unknown_columns", unknown_columns=["interface_name"]),
        ImportFileTooLarge(10),
        ImportRowLimitExceeded(1),
        ImportValidationFailed(
            total_rows=1,
            importable_rows=0,
            error_count=1,
            errors=[
                {"row": 2, "field": "cluster", "reason": "not_found", "related_row": None}
            ],
            errors_truncated=False,
        ),
    ]
    good = xlsx_bytes(BM_HEADERS, _sample_rows(BM_HEADERS))
    for exc in app_errors:
        with monkeypatch.context() as patch:
            patch.setattr("openpyxl.load_workbook", _raising(exc))
            with pytest.raises(type(exc)) as perr:
                resource_import.parse_records(good, "xlsx")
        assert perr.value is exc, "the very same AppError instance must propagate"
        assert perr.value.status_code == exc.status_code
        assert perr.value.details == exc.details
        assert perr.value.__cause__ is None

    # CSV 路径同样保真
    with monkeypatch.context() as patch:
        patch.setattr(
            resource_import.csv, "reader", _raising(ImportInvalidFile("empty_file"))
        )
        with pytest.raises(ImportInvalidFile) as csv_app_error:
            resource_import.parse_records(b"cluster\n", "csv")
    assert csv_app_error.value.details == {"reason": "empty_file"}


@pytest.mark.parametrize("exc_type", [MemoryError, RecursionError])
def test_td_f037_u6_memory_and_recursion_errors_propagate(monkeypatch, exc_type):
    """TD-F037-U6（U-d / R-6）：MemoryError / RecursionError 不得吞成 422。"""
    good = xlsx_bytes(BM_HEADERS, _sample_rows(BM_HEADERS))
    with monkeypatch.context() as patch:
        patch.setattr("openpyxl.load_workbook", _raising(exc_type()))
        with pytest.raises(exc_type):
            resource_import.parse_records(good, "xlsx")

    with monkeypatch.context() as patch:
        patch.setattr(resource_import.csv, "reader", _raising(exc_type()))
        with pytest.raises(exc_type):
            resource_import.parse_records(b"cluster,hostname\nA,b\n", "csv")


@pytest.mark.parametrize("exc_type", [KeyboardInterrupt, SystemExit])
def test_td_f037_u7_interrupt_semantics_unchanged(monkeypatch, exc_type):
    """TD-F037-U7：KeyboardInterrupt / SystemExit 语义不变（BaseException 不捕）。"""
    good = xlsx_bytes(BM_HEADERS, _sample_rows(BM_HEADERS))
    exc = exc_type(3) if exc_type is SystemExit else exc_type()
    with monkeypatch.context() as patch:
        patch.setattr("openpyxl.load_workbook", _raising(exc))
        with pytest.raises(exc_type) as perr:
            resource_import.parse_records(good, "xlsx")
    if exc_type is SystemExit:
        assert perr.value.code == 3


@pytest.mark.parametrize(
    "exc_type",
    [AttributeError, TypeError, IndexError, AssertionError, OverflowError, SyntaxError],
)
def test_td_f037_u8_u16_non_whitelisted_errors_propagate(monkeypatch, exc_type):
    """TD-F037-U8 + §15.8.4：未白名单异常（含裸 SyntaxError）原样传播。"""
    good = xlsx_bytes(BM_HEADERS, _sample_rows(BM_HEADERS))
    exc = exc_type("injected")
    with monkeypatch.context() as patch:
        patch.setattr("openpyxl.load_workbook", _raising(exc))
        with pytest.raises(exc_type) as perr:
            resource_import.parse_records(good, "xlsx")
    assert not isinstance(perr.value, ImportInvalidFile)

    with monkeypatch.context() as patch:
        patch.setattr(resource_import.csv, "reader", _raising(exc_type("injected")))
        with pytest.raises(exc_type):
            resource_import.parse_records(b"cluster,hostname\nA,b\n", "csv")


def test_td_f037_u9_decision_order_and_materialised_result():
    """TD-F037-U9（含 TD-I-8 修订）：判定顺序与「全量物化列表」语义不变。"""
    with pytest.raises(ImportInvalidFile) as csv_empty:
        resource_import.parse_records(b"", "csv")
    assert csv_empty.value.details == {"reason": "empty_file"}

    with pytest.raises(ImportInvalidFile) as xlsx_empty:
        resource_import.parse_records(b"", "xlsx")
    assert xlsx_empty.value.details == {"reason": "empty_file"}

    with pytest.raises(ImportInvalidFile) as ole2:
        resource_import.parse_records(OLE2_MAGIC, "xlsx")
    assert ole2.value.details == {"reason": "unsupported_format"}

    # 零成员 zip：magic 为 PK\x05\x06，被 PK 闸门拒绝（TD-I-8 修订：不是 empty_file）
    with pytest.raises(ImportInvalidFile) as zero_member:
        resource_import.parse_records(zero_member_zip(), "xlsx")
    assert zero_member.value.details == {"reason": "unsupported_format"}

    with pytest.raises(ImportInvalidFile) as bad_zip:
        resource_import.parse_records(b"PK\x03\x04not-a-zip", "xlsx")
    assert bad_zip.value.details == {"reason": "unreadable"}
    assert isinstance(bad_zip.value.__cause__, zipfile.BadZipFile)

    # empty_file（合法容器、workbook.worksheets 为空）：先于 iter_rows
    with pytest.raises(ImportInvalidFile) as no_sheet:
        resource_import.parse_records(xlsx_without_sheet_part(BM_HEADERS), "xlsx")
    assert no_sheet.value.details == {"reason": "empty_file"}
    assert no_sheet.value.__cause__ is None

    records = resource_import.parse_records(
        xlsx_bytes(BM_HEADERS, _sample_rows(BM_HEADERS)), "xlsx"
    )
    assert isinstance(records, list) and records
    assert all(isinstance(row, list) for row in records)
    assert all(isinstance(cell, str) for row in records for cell in row)
    assert records[0] == BM_HEADERS


def test_td_f037_u10_close_lifecycle_and_no_override(monkeypatch):
    """TD-F037-U10（R-2）：close 恰好一次、异常被吞、不覆盖已判定 422。"""
    import xml.etree.ElementTree as ElementTree

    calls: list[object] = []
    real_close = resource_import._close_workbook

    def spy(workbook):
        calls.append(workbook)
        return real_close(workbook)

    class FakeSheet:
        def iter_rows(self, values_only: bool = True):
            raise ElementTree.ParseError("mismatched tag: line 1, column 9")

    class FakeWorkbook:
        worksheets = [FakeSheet()]
        close_attempts = 0

        def close(self):
            FakeWorkbook.close_attempts += 1
            raise RuntimeError("close boom")

    good = xlsx_bytes(BM_HEADERS, _sample_rows(BM_HEADERS))

    # (a) 正常路径：close 恰 1 次
    monkeypatch.setattr(resource_import, "_close_workbook", spy)
    resource_import.parse_records(good, "xlsx")
    assert len(calls) == 1

    # (b) 解析失败 + close 抛错：仍是 422 unreadable，且 close 恰 1 次
    calls.clear()
    with monkeypatch.context() as patch:
        patch.setattr("openpyxl.load_workbook", lambda *_a, **_k: FakeWorkbook())
        with pytest.raises(ImportInvalidFile) as perr:
            resource_import.parse_records(good, "xlsx")
    assert perr.value.details == {"reason": "unreadable"}
    assert isinstance(perr.value.__cause__, ElementTree.ParseError)
    assert not isinstance(perr.value.__cause__, RuntimeError)
    assert FakeWorkbook.close_attempts == 1
    assert len(calls) == 1

    # (c) load_workbook 自身失败：workbook is None ⇒ close 0 次
    calls.clear()
    with monkeypatch.context() as patch:
        patch.setattr(
            "openpyxl.load_workbook", _raising(zipfile.BadZipFile("bad zip"))
        )
        with pytest.raises(ImportInvalidFile) as load_failure:
            resource_import.parse_records(good, "xlsx")
    assert load_failure.value.details == {"reason": "unreadable"}
    assert calls == []


def test_td_f037_u11_cell_text_errors_stay_outside_the_conversion_scope(monkeypatch):
    """TD-F037-U11（R-9）：转换作用域最小性；表头/列类 reason 不受影响。"""
    good = xlsx_bytes(BM_HEADERS, _sample_rows(BM_HEADERS))

    with monkeypatch.context() as patch:
        patch.setattr(
            resource_import,
            "_cell_text",
            _raising(TypeError("cell normalisation bug")),
        )
        with pytest.raises(TypeError):
            resource_import.parse_records(b"cluster,hostname\nA,b\n", "csv")
        with pytest.raises(TypeError):
            resource_import.parse_records(good, "xlsx")

    with pytest.raises(ImportInvalidFile) as unknown:
        resource_import._resolve_template(BM_HEADERS + ["interface_name"])
    assert unknown.value.details["reason"] == "unknown_columns"
    with pytest.raises(ImportInvalidFile) as missing:
        resource_import._resolve_template(BM_HEADERS[:-1])
    assert missing.value.details["reason"] == "missing_columns"
    swapped = list(BM_HEADERS)
    swapped[0], swapped[1] = swapped[1], swapped[0]
    with pytest.raises(ImportInvalidFile) as order:
        resource_import._resolve_template(swapped)
    assert order.value.details["reason"] == "column_order"


def test_td_f037_u12_parse_logging_is_content_free(caplog):
    """TD-F037-U12（A18 前半）：日志只含 format/bytes/异常类名。"""
    caplog.set_level(logging.INFO, logger="csm.imports")

    with pytest.raises(ImportInvalidFile):
        resource_import.parse_records(oversized_field_csv(BM_HEADERS, marker=MARKER), "csv")
    with pytest.raises(ImportInvalidFile):
        resource_import.parse_records(corrupt_sheet_tail(BM_HEADERS, marker=MARKER), "xlsx")
    with pytest.raises(ImportInvalidFile):
        resource_import.parse_records(corrupt_manifest(BM_HEADERS), "xlsx")

    records = [r for r in caplog.records if r.name == "csm.imports"]
    rejection_records = [
        r for r in records if "import parse rejected" in str(r.msg)
    ]
    for record in records:
        assert record.exc_info is None, record.msg
        assert MARKER not in str(record.msg)
    for record in rejection_records:
        assert isinstance(record.args, tuple), record.args
        fmt, size, error_name = record.args
        assert fmt in {"csv", "xlsx"}, fmt
        assert isinstance(size, int) and size > 0, size
        assert re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", str(error_name)), error_name

    text_dump = caplog.text
    for forbidden in (
        MARKER,
        "<worksheet",
        "mismatched tag",
        "field larger than field limit",
        "Traceback",
        PASTED_FILE_NAME,
    ):
        assert forbidden not in text_dump, f"log leaked {forbidden!r}"


def test_td_f037_u13_missing_workbook_part_converges(monkeypatch):
    """TD-F037-U13（§15.10.7-1）：缺 workbook 部件的 zip ⇒ 422 + __cause__ OSError。"""
    raw = manifest_without_workbook_part()
    assert raw.startswith(b"PK\x03\x04")

    with pytest.raises(OSError) as excinfo:
        load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    assert "no valid workbook part" in str(excinfo.value), str(excinfo.value)

    with pytest.raises(ImportInvalidFile) as perr:
        resource_import.parse_records(raw, "xlsx")
    assert perr.value.details == {"reason": "unreadable"}
    assert type(perr.value.__cause__) is OSError

    # OSError 未进入模块级 `_PARSE_FAILURES`（只局部生效）：CSV 路径必须传播
    assert OSError not in resource_import._PARSE_FAILURES


def test_td_f037_u14_oserror_scope_is_minimal(monkeypatch):
    """TD-F037-U14（§15.10.7-3 / R-12）：OSError 仅在 _parse_xlsx 作用域内收敛。"""
    with monkeypatch.context() as patch:
        patch.setattr(resource_import.csv, "reader", _raising(OSError("csv boom")))
        with pytest.raises(OSError):
            resource_import.parse_records(b"cluster,hostname\nA,b\n", "csv")

    good = xlsx_bytes(BM_HEADERS, _sample_rows(BM_HEADERS))
    with monkeypatch.context() as patch:
        patch.setattr(
            resource_import, "_cell_text", _raising(OSError("normalisation boom"))
        )
        with pytest.raises(OSError):
            resource_import.parse_records(good, "xlsx")


# =============================================================================
# I 层 — 真实 HTTP + 真实 PostgreSQL
# =============================================================================

PARSE_FAILURE_CASES = [
    pytest.param(RES_VALIDATE, "csv", id="I1-resources-validate-csv"),
    pytest.param(RES_VALIDATE, "xlsx", id="I2-resources-validate-xlsx"),
    pytest.param(RES_WRITE, "csv", id="I3-resources-write-csv"),
    pytest.param(RES_WRITE, "xlsx", id="I4-resources-write-xlsx"),
    pytest.param(UNI_VALIDATE, "csv", id="I5-unified-validate-csv"),
    pytest.param(UNI_VALIDATE, "xlsx", id="I6-unified-validate-xlsx"),
    pytest.param(UNI_WRITE, "csv", id="I7-unified-write-csv"),
    pytest.param(UNI_WRITE, "xlsx", id="I8-unified-write-xlsx"),
]


@pytest.mark.parametrize("path, fmt", PARSE_FAILURE_CASES)
def test_td_f037_i1_i8_parse_failure_is_422_at_every_endpoint(
    role_client_factory, path, fmt
):
    """TD-F037-I1…I8（A1–A8 + A20 前置）：4 端点 × 2 输入全 422 unreadable。"""
    editor, _user, _password = role_client_factory("EDITOR")
    headers = headers_for(path)
    payload = (
        oversized_field_csv(headers) if fmt == "csv" else corrupt_sheet_tail(headers)
    )

    resp = post(editor, path, payload, fmt)
    assert_unreadable(resp, f"{path} {fmt}")
    assert set(resp.json()["error"]) == {"code", "message", "details"}
    again = post(editor, path, payload, fmt)
    assert again.json() == resp.json(), "same bytes must yield the same body (A20)"


def test_td_f037_i9_rejected_write_leaves_no_residue(role_client_factory, auth_session_factory):
    """TD-F037-I9（A9）：4 条写入路径零残留；audit_logs 不变即覆盖 Activity。"""
    editor, _user, _password = role_client_factory("EDITOR")

    with auth_session_factory() as session:
        before = counts(session)

    responses = []
    for path in WRITE_PATHS:
        headers = headers_for(path)
        responses.append(post(editor, path, oversized_field_csv(headers), "csv"))
        responses.append(post(editor, path, corrupt_sheet_tail(headers), "xlsx"))

    for resp in responses:
        assert_unreadable(resp)
        assert resp.status_code != 201

    with auth_session_factory() as session:
        after = counts(session)
    assert after == before, f"rejected writes must not persist rows: {before} -> {after}"


def test_td_f037_i10_auth_and_permission_precede_parsing(auth_app, role_client_factory):
    """TD-F037-I10（A10 / B10）：401 与 403 均先于解析，不得出现 422 / 500。"""
    unauth = TestClient(auth_app)
    unauth.__enter__()
    try:
        for path in ALL_PATHS:
            payload = oversized_field_csv(headers_for(path))
            resp = post(unauth, path, payload)
            assert resp.status_code == 401, f"{path}: {resp.status_code} {resp.text[:300]}"
            assert resp.json()["error"]["code"] == "AUTH_REQUIRED"
    finally:
        unauth.__exit__(None, None, None)

    viewer, _user, _password = role_client_factory("VIEWER")
    for path in ALL_PATHS:
        payload = oversized_field_csv(headers_for(path))
        resp = post(viewer, path, payload)
        assert resp.status_code == 403, f"{path}: {resp.status_code} {resp.text[:300]}"
        assert resp.json()["error"] == {
            "code": "FORBIDDEN",
            "message": "Forbidden",
            "details": {"reason": "permission"},
        }


def test_td_f037_i11_byte_limit_still_wins(role_client_factory, monkeypatch):
    """TD-F037-I11（A11 / B4）：413 不得被 422 取代；输入 ① 必须越过 413。"""
    editor, _user, _password = role_client_factory("EDITOR")

    with monkeypatch.context() as patch:
        patch.setattr(
            imports_api,
            "get_settings",
            lambda: Settings(_env_file=None, import_max_file_bytes=10),
        )
        for path in ALL_PATHS:
            payload = oversized_field_csv(headers_for(path))
            resp = post(editor, path, payload)
            assert resp.status_code == 413, f"{path}: {resp.status_code} {resp.text[:300]}"
            assert resp.json()["error"] == {
                "code": "IMPORT_FILE_TOO_LARGE",
                "message": "Import file too large",
                "details": {"limit_bytes": 10},
            }

    for path in ALL_PATHS:
        payload = oversized_field_csv(headers_for(path))
        resp = post(editor, path, payload)
        assert_unreadable(resp, f"default limits {path}")


def test_td_f037_i12_empty_file_reason_not_converged(role_client_factory):
    """TD-F037-I12（A12 / B1）：空字节与仅表头仍是 empty_file。"""
    editor, _user, _password = role_client_factory("EDITOR")
    for path in ALL_PATHS:
        headers = headers_for(path)
        assert_invalid_file(post(editor, path, b""), "empty_file", f"{path} empty bytes")
        assert_invalid_file(
            post(editor, path, header_only_csv(headers)),
            "empty_file",
            f"{path} header only",
        )


def test_td_f037_i13_unsupported_format_not_converged(role_client_factory):
    """TD-F037-I13（A13 / B2）：OLE2 仍 unsupported_format；PK 非 zip 仍 unreadable。"""
    editor, _user, _password = role_client_factory("EDITOR")
    for path in ALL_PATHS:
        ole2 = post(editor, path, OLE2_MAGIC, "xlsx")
        assert_invalid_file(ole2, "unsupported_format", f"{path} ole2")

        broken = post(editor, path, b"PK\x03\x04not-a-zip", "xlsx")
        assert_unreadable(broken, f"{path} pk-not-zip")


def test_td_f037_i14_header_and_column_reasons_unchanged(role_client_factory, monkeypatch):
    """TD-F037-I14（A14 / B3）：表头/列类 reason 与「行数先于表头」顺序不变。"""
    editor, _user, _password = role_client_factory("EDITOR")
    row = {"cluster": "F37A", "hostname": "f37-x", "resource_type": "BARE_METAL"}

    unknown = post(editor, RES_VALIDATE, csv_bytes(BM_HEADERS + ["interface_name"], [row]))
    assert_invalid_file(unknown, "unknown_columns", "resources unknown column")
    assert unknown.json()["error"]["details"]["unknown_columns"] == ["interface_name"]

    missing = post(editor, RES_VALIDATE, csv_bytes(BM_HEADERS[:-1], []))
    assert_invalid_file(missing, "missing_columns", "resources missing column")
    assert missing.json()["error"]["details"]["missing_columns"] == ["gpu_count"]

    swapped = list(BM_HEADERS)
    swapped[0], swapped[1] = swapped[1], swapped[0]
    order = post(editor, RES_VALIDATE, csv_bytes(swapped, [row]))
    assert_invalid_file(order, "column_order", "resources column order")

    header_missing = post(editor, RES_VALIDATE, b",,,,,,,,,,,,,\n" + b"F37A,,,,,,,,,,,,,\n")
    assert_invalid_file(header_missing, "header_missing", "resources header missing")

    unified_unknown = post(
        editor, UNI_VALIDATE, csv_bytes(W26_HEADERS + ["extra"], [])
    )
    assert_invalid_file(unified_unknown, "unknown_columns", "unified unknown column")
    unified_missing = post(editor, UNI_VALIDATE, csv_bytes(W26_HEADERS[:-1], []))
    assert_invalid_file(unified_missing, "missing_columns", "unified missing column")

    with monkeypatch.context() as patch:
        patch.setattr(
            imports_api, "get_settings", lambda: Settings(_env_file=None, import_max_rows=1)
        )
        two_rows = [row, {**row, "hostname": "f37-y"}]
        over = post(editor, RES_VALIDATE, csv_bytes(BM_HEADERS + ["interface_name"], two_rows))
    assert over.status_code == 413, over.text[:300]
    assert over.json()["error"] == {
        "code": "IMPORT_ROW_LIMIT_EXCEEDED",
        "message": "Import row limit exceeded",
        "details": {"limit_rows": 1},
    }


def test_td_f037_i15_row_limit_413_without_partial_import(
    role_client_factory, auth_session_factory
):
    """TD-F037-I15（A15 / T111）：5001 行 → 413，零写入。"""
    editor, _user, _password = role_client_factory("EDITOR")

    def _many_rows(headers: list[str], count: int) -> bytes:
        row = ["F37A", "h"] + [""] * (len(headers) - 2)
        return csv_bytes(headers, [dict(zip(headers, row)) for _ in range(count)])

    with auth_session_factory() as session:
        before = counts(session)
    for path in ALL_PATHS:
        resp = post(editor, path, _many_rows(headers_for(path), 5001))
        assert resp.status_code == 413, f"{path}: {resp.status_code} {resp.text[:300]}"
        assert resp.json()["error"] == {
            "code": "IMPORT_ROW_LIMIT_EXCEEDED",
            "message": "Import row limit exceeded",
            "details": {"limit_rows": 5000},
        }
    with auth_session_factory() as session:
        after = counts(session)
    assert after == before


def test_td_f037_i16_row_errors_shape_and_one_based_rows(
    role_client_factory, auth_session_factory
):
    """TD-F037-I16（A16 / B5）：行级 RowError 形状、1-based 行号与 422 写入形状不变。"""
    editor, _user, _password = role_client_factory("EDITOR")
    cluster = editor.post("/api/v1/clusters", json={"code": "C37", "name": "Cluster"})
    assert cluster.status_code == 201, cluster.text

    rows = [
        {"cluster": "C37", "hostname": "ok-host", "resource_type": "BARE_METAL"},
        {"cluster": "C37", "hostname": "bad name", "resource_type": "BARE_METAL"},
    ]
    data = csv_bytes(BM_HEADERS, rows)

    validate = post(editor, RES_VALIDATE, data)
    assert validate.status_code == 200, validate.text[:400]
    payload = validate.json()
    assert payload["total_rows"] == 2
    assert payload["error_count"] == 1
    assert payload["importable_rows"] == 1
    error = payload["errors"][0]
    assert set(error) == {"row", "field", "reason", "related_row"}
    assert error["row"] == 3 and error["field"] == "hostname"

    with auth_session_factory() as session:
        before = counts(session)
    write = post(editor, RES_WRITE, data)
    assert write.status_code == 422, write.text[:400]
    body = write.json()["error"]
    assert body["code"] == "IMPORT_VALIDATION_FAILED"
    assert body["message"] == "Import validation failed"
    assert set(body["details"]) == {
        "total_rows",
        "importable_rows",
        "error_count",
        "errors",
        "errors_truncated",
    }
    assert body["details"]["errors"][0]["row"] == 3
    with auth_session_factory() as session:
        after = counts(session)
    assert after == before, "IMPORT_VALIDATION_FAILED must not write a partial batch"

    unified_rows = [{"cluster": "NOPE"}]
    unified = post(editor, UNI_VALIDATE, csv_bytes(W26_HEADERS, unified_rows))
    assert unified.status_code == 200, unified.text[:400]
    unified_error = unified.json()["errors"][0]
    assert set(unified_error) == {"row", "field", "reason", "related_row"}
    assert unified_error["row"] == 2


def test_td_f037_i17_dual_carrier_equivalence(role_client_factory):
    """TD-F037-I17（A17 / B9 / T114）：同字节两载体响应逐字段相同。"""
    editor, _user, _password = role_client_factory("EDITOR")
    for path in ALL_PATHS:
        payload = oversized_field_csv(headers_for(path))
        file_resp = post(editor, path, payload, "csv", carrier="file")
        text_resp = post(editor, path, payload, "csv", carrier="text")
        assert file_resp.status_code == text_resp.status_code, path
        assert file_resp.json() == text_resp.json(), path
        assert_unreadable(file_resp, f"{path} file carrier")
        assert_unreadable(text_resp, f"{path} text carrier")

    # A17 补充：服务层异常 details 与 HTTP 响应 details 一致（单一定义来源）
    with pytest.raises(ImportInvalidFile) as perr:
        resource_import.parse_records(oversized_field_csv(BM_HEADERS), "csv")
    assert perr.value.details == {"reason": "unreadable"}


def test_td_f037_i18_response_body_is_content_free(role_client_factory):
    """TD-F037-I18（A18 后半）：响应体不含上传内容、XML 片段或堆栈。"""
    editor, _user, _password = role_client_factory("EDITOR")
    for path in ALL_PATHS:
        headers = headers_for(path)
        payloads = {
            "csv": oversized_field_csv(headers, marker=MARKER),
            "xlsx": corrupt_sheet_tail(headers, marker=MARKER),
        }
        for fmt, payload in payloads.items():
            resp = post(editor, path, payload, fmt)
            assert_unreadable(resp, f"{path} {fmt}")
            body = resp.text
            for forbidden in (
                MARKER,
                "<worksheet",
                "mismatched tag",
                "field larger than field limit",
                "Traceback",
                ".py",
            ):
                assert forbidden not in body, f"{path} {fmt} leaked {forbidden!r}"


def test_td_f037_i19_internal_errors_stay_500(auth_app, make_user, monkeypatch):
    """TD-F037-I19（A19 补充 / B7–B8）：未白名单异常仍为 500 INTERNAL_ERROR。"""
    user, password = make_user(role="EDITOR")
    client = TestClient(auth_app, raise_server_exceptions=False)
    client.__enter__()
    try:
        login = client.post(
            "/api/v1/auth/login",
            json={"username": user.username, "password": password},
        )
        assert login.status_code == 200, login.text

        with monkeypatch.context() as patch:
            patch.setattr(
                resource_import, "parse_records", _raising(MemoryError("oom"))
            )
            for path in ALL_PATHS:
                payload = oversized_field_csv(headers_for(path))
                resp = post(client, path, payload)
                assert resp.status_code == 500, f"{path}: {resp.status_code} {resp.text[:300]}"
                assert resp.json() == {
                    "error": {
                        "code": "INTERNAL_ERROR",
                        "message": "Internal server error",
                        "details": {},
                    }
                }
    finally:
        client.__exit__(None, None, None)


def test_td_f037_i20_repeated_validate_is_idempotent_and_read_only(
    role_client_factory, auth_session_factory
):
    """TD-F037-I20（A20 / T116）：重复 validate 一致、零写入、零 Audit。"""
    editor, _user, _password = role_client_factory("EDITOR")
    with auth_session_factory() as session:
        before = counts(session)

    for path in (RES_VALIDATE, UNI_VALIDATE):
        payload = oversized_field_csv(headers_for(path))
        first = post(editor, path, payload)
        second = post(editor, path, payload)
        assert_unreadable(first, f"{path} first")
        assert first.json() == second.json(), path

    with auth_session_factory() as session:
        after = counts(session)
    assert after == before


def test_td_f037_i21_no_file_descriptor_leak(role_client_factory):
    """TD-F037-I21（R-5）：20 次损坏 xlsx 提交后 fd 数不增长。"""
    fd_dir = Path("/proc/self/fd")
    if not fd_dir.is_dir():
        pytest.skip("/proc/self/fd unavailable on this platform")
    editor, _user, _password = role_client_factory("EDITOR")
    payload = corrupt_sheet_tail(BM_HEADERS)

    before = len(os.listdir(fd_dir))
    for _index in range(20):
        assert_unreadable(post(editor, RES_VALIDATE, payload, "xlsx"))
    after = len(os.listdir(fd_dir))
    assert after <= before, f"fd leak: {before} -> {after}"


def test_td_f037_i22_missing_workbook_part_is_422_at_every_endpoint(role_client_factory):
    """TD-F037-I22（§15.10.7-4 / BE-01）：4 端点 × 该字节 → 422，无 500/413/2xx。"""
    editor, _user, _password = role_client_factory("EDITOR")
    payload = manifest_without_workbook_part()
    for path in ALL_PATHS:
        resp = post(editor, path, payload, "xlsx")
        assert_unreadable(resp, f"{path} no-workbook-part")
        assert resp.status_code not in (413, 500)


def test_td_f037_i23_corrupt_manifest_is_422_at_every_endpoint(role_client_factory):
    """architecture §15.8.2（输入 ③）：4 端点 → 422 unreadable。"""
    editor, _user, _password = role_client_factory("EDITOR")
    payload = corrupt_manifest(BM_HEADERS)
    for path in ALL_PATHS:
        resp = post(editor, path, payload, "xlsx")
        assert_unreadable(resp, f"{path} corrupt-manifest")
        assert resp.status_code not in (413, 500)


# =============================================================================
# S 层 — 源码级静态守卫（`git diff` 类守卫见 test-report §2）
# =============================================================================

SOURCE_PATH = Path(resource_import.__file__)


def _source() -> str:
    return SOURCE_PATH.read_text(encoding="utf-8")


def _function_node(tree: ast.Module, name: str) -> ast.FunctionDef:
    for node in ast.walk(tree):
        if isinstance(node, ast.FunctionDef) and node.name == name:
            return node
    raise AssertionError(f"function {name} not found")


def test_td_f037_s2_caps_and_csv_strictness_unchanged():
    """TD-F037-S2（A21 / T126）：字段上限、严格引号与三个上限常量零变更。"""
    source = _source()
    assert "field_size_limit" not in source
    assert "strict=True" not in source
    assert "csv.QUOTE" not in source

    settings = Settings(_env_file=None)
    assert settings.import_max_file_bytes == 5_242_880
    assert settings.import_max_rows == 5000
    assert settings.import_error_preview_limit == 1000


def test_td_f037_s3_logging_shape_and_classification_boundary():
    """TD-F037-S3（A18 / A21）：无 exception/error 日志、无 filename/exc_info。"""
    source = _source()
    assert "logger.exception" not in source
    assert "logger.error" not in source
    assert "exc_info" not in source
    assert "filename" not in source
    assert "except SyntaxError" not in source

    tree = ast.parse(source)
    generic_handlers = [
        node.lineno
        for node in ast.walk(tree)
        if isinstance(node, ast.ExceptHandler)
        and node.type is not None
        and ast.unparse(node.type) == "Exception"
    ]
    # 只剩两处既有边界：_close_workbook 清理 + write_import 事务回滚（零 diff）
    assert len(generic_handlers) == 2, generic_handlers
    close_workbook = _function_node(tree, "_close_workbook")
    write_import = _function_node(tree, "write_import")
    for lineno in generic_handlers:
        assert close_workbook.lineno <= lineno <= close_workbook.end_lineno or (
            write_import.lineno <= lineno <= write_import.end_lineno
        ), lineno

    converter = _function_node(tree, "_invalid_file_on_parse_failure")
    handler_types = []
    for node in ast.walk(converter):
        if isinstance(node, ast.ExceptHandler):
            handler_types.append(ast.unparse(node.type) if node.type else "bare")
    assert handler_types == ["AppError", "failures"], handler_types


def test_td_f037_s5_only_signed_reasons_and_frozen_templates():
    """TD-F037-S5（A21 / T126）：reason 集合零新增；BM 14 / VM 10 列序冻结。"""
    source = _source()
    literal_reasons = set(re.findall(r'ImportInvalidFile\(\s*"([a-z_]+)"', source))
    assert literal_reasons <= SIGNED_REASONS, literal_reasons
    assert "unreadable" in literal_reasons

    templates = {item["resource_type"]: item for item in resource_import.templates_payload()}
    assert [column["name"] for column in templates["BARE_METAL"]["columns"]] == BM_HEADERS
    assert len(templates["VIRTUAL_MACHINE"]["columns"]) == 10
    frozen = {template.resource_type: template.names for template in resource_import.TEMPLATES}
    assert frozen["BARE_METAL"] == BM_HEADERS
    assert list(frozen) == ["BARE_METAL", "VIRTUAL_MACHINE"]


def test_td_f037_s6_validate_path_has_no_locks_or_commits():
    """TD-F037-S6（A20 补充）：解析/预校验路径无锁、无 SAVEPOINT、无 commit。"""
    import inspect

    for name in ("parse_records", "parse_and_validate", "validate_import", "_parse_csv", "_parse_xlsx"):
        source = inspect.getsource(getattr(resource_import, name))
        for forbidden in ("for_update", "begin_nested", ".commit(", "open(", "Path("):
            assert forbidden not in source, f"{name} must not reference {forbidden!r}"


def test_td_f037_s7_oserror_is_local_to_parse_xlsx():
    """TD-F037-S7（§15.10.7-5 / R-12）：OSError 仅局部；with 块内无真实 I/O。"""
    source = _source()
    tree = ast.parse(source)

    parser = _function_node(tree, "_parse_xlsx")
    start, end = parser.lineno, parser.end_lineno
    for index, line in enumerate(source.splitlines(), start=1):
        if "OSError" in line:
            assert start <= index <= end, f"OSError referenced outside _parse_xlsx: {line}"

    # `_PARSE_FAILURES` 未被放宽
    assert len(resource_import._PARSE_FAILURES) == 8
    assert OSError not in resource_import._PARSE_FAILURES

    # 无独立 `except OSError` 分支
    for node in ast.walk(parser):
        if isinstance(node, ast.ExceptHandler) and node.type is not None:
            assert "OSError" not in ast.unparse(node.type)

    # 转换作用域内只有对内存字节的第三方调用（L7′）
    forbidden_calls = {"open", "Path", "NamedTemporaryFile", "TemporaryFile"}
    with_blocks = [node for node in ast.walk(parser) if isinstance(node, ast.With)]
    assert with_blocks, "_parse_xlsx must keep its conversion scope"
    for block in with_blocks:
        for node in ast.walk(block):
            if isinstance(node, ast.Call):
                name = ast.unparse(node.func)
                tail = name.split(".")[-1]
                assert tail not in forbidden_calls, f"real I/O in conversion scope: {name}"


def test_td_f037_s8_lxml_stays_an_optional_dependency():
    """architecture §15.8.5：lxml 不落依赖清单，仅 resource_import.py 条件导入。"""
    backend_dir = SOURCE_PATH.parents[2]
    pyproject = (backend_dir / "pyproject.toml").read_text(encoding="utf-8")
    assert "lxml" not in pyproject

    hits = sorted(
        path.relative_to(backend_dir).as_posix()
        for path in (backend_dir / "app").rglob("*.py")
        if "lxml" in path.read_text(encoding="utf-8")
    )
    assert hits == ["app/services/resource_import.py"], hits
