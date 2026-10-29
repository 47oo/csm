"""F038 导入解压 / 解析内存边界加固 — 独立 TEST 执行（U / I / M / S）。

覆盖 `docs/features/F038/test-report.md` §1 `TD-F038-1` 的 U1–U6 / I1–I15 /
M1–M6 / S1–S6，并按 `docs/features/F038/architecture.md` §13 的全部裁定执行：

* §13.1 毒值探针 P-1..P-4（含双负向对照与错误路径）；
* §13.2 阈值口径 T-1 / T-1b / T-2 / T-3 / T-4 / T-5 / T-6（公式原样，**不放宽**）；
  **repair round 1 后 CSV 族按 §14.2 执行**：T-1b′ / T-1c / T-2′ 改用实测地板
  探针 `F(x)` 剥离后的工作集 `work(x) = max(peak − F(x), 0)`，乘数/余量零变更；
  T-1（xlsx）/ T-3 / T-4 / T-5 / T-6 不受影响；
* §13.3 ON-1：尾部**解析器异常** → 422 `unreadable`；尾部**行级错误/超宽行** → 413；
* §13.4 S-W2 期望 = **422 `unknown_columns`**（不是 413）；
* §13.5 E-1/E-2/E-3：`parse_records` 两参数兼容 + 全模块调用点必须带两关键字。

分层（TD-2）：

* **U** — 直调 `app.services.resource_import` / `resource_interface_ip_import`
  的私有函数；U4 毒值探针需要真实 session 解析 cluster/network，故使用真实库；
* **I** — `auth_app`（真实 `create_app`）+ 真实 PostgreSQL（`TEST_DATABASE_URL`，
  DB-D10 双守卫），4 端点 × 输入矩阵，**基线语义路径 vs 生产路径**深比较；
* **M** — `tracemalloc` 峰值 + `perf_counter`，基线来自 `git show <start_commit>:…`
  写入 `/tmp`（**不入库**）后用 `importlib` 载入；
* **S** — 源码 AST / grep / 只读 `git diff`。

全部样例在内存中构造（`io.BytesIO` / `zipfile` / `csv.writer`），不落盘、不读写
被跟踪的 artifacts。本模块不导入 `app.core.access_scope` 符号，故按 F030 §9.1
不需要 `pytest.mark.real_scope`（由 `test_f030_test_harness.py` 的 AST 门禁复核）。
"""

from __future__ import annotations

import ast
import csv
import gc
import importlib.util
import io
import os
import re
import subprocess
import sys
import time
import tracemalloc
import zipfile
from datetime import date, datetime, time as dtime, timedelta
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from app.api.v1 import imports as imports_api
from app.core.config import Settings
from app.core.errors import (
    ImportInvalidFile,
    ImportRowLimitExceeded,
)
from app.services import resource_import
from app.services import resource_interface_ip_import as unified

from integration.f015_helpers import (  # noqa: F401  (reused fixture)
    create_cluster,
    create_network,
    role_client_factory,
)

# --- constants ---------------------------------------------------------------

RES_VALIDATE = "/api/v1/imports/resources/validate"
RES_WRITE = "/api/v1/imports/resources"
UNI_VALIDATE = "/api/v1/imports/resource-interfaces-ips/validate"
UNI_WRITE = "/api/v1/imports/resource-interfaces-ips"
ALL_PATHS = (RES_VALIDATE, RES_WRITE, UNI_VALIDATE, UNI_WRITE)
WRITE_PATHS = (RES_WRITE, UNI_WRITE)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# F014 BM 14 列 / F023 v2 26 列冻结列序（Contract §6；零变更由 S1 复核）。
BM_HEADERS = [
    "cluster",
    "hostname",
    "resource_type",
    "role",
    "status",
    "manufacturer",
    "model",
    "serial_number",
    "cpu_model",
    "cpu_sockets",
    "cpu_cores",
    "memory_mb",
    "gpu_model",
    "gpu_count",
]
W26_HEADERS = [
    "cluster",
    "hostname",
    "resource_type",
    "role",
    "status",
    "host",
    "vcpu",
    "memory_mb",
    "disk_gb",
    "hypervisor",
    "manufacturer",
    "model",
    "serial_number",
    "cpu_model",
    "cpu_sockets",
    "cpu_cores",
    "gpu_model",
    "gpu_count",
    "name",
    "interface_type",
    "mac_address",
    "speed_mbps",
    "mtu",
    "enabled",
    "network",
    "ip",
]

# 载体形状常量（与 F034 `importCarrier.ts` 同值，见 test_f034_csv_text_carrier.py）。
PASTED_FILE_NAME = "csm-pasted-import.csv"
PASTED_MIME = "text/csv;charset=utf-8"
FILE_NAME = "f038-file-carrier.csv"
FILE_MIME = "text/csv"

MARKER = "F038-MARKER-7c2f1d9b4a6e5c30"

COUNT_TABLES = (
    "clusters",
    "resources",
    "bare_metal_details",
    "virtual_machine_details",
    "interfaces",
    "ip_addresses",
    "audit_logs",
)

CAP_BYTES = 5_242_880
CAP_ROWS = 5000
CAP_PREVIEW = 1000

CLUSTER = "F038"
NETWORK = "mgmt"
NETWORK_CIDR = "10.38.0.0/24"

# start_commit（F038 起点的语义基线）；`/tmp` 基线模块由 git show 生成、不入库。
START_COMMIT = "8e609ca"
BASELINE_PATH = os.environ.get(
    "F038_BASELINE_MODULE", "/tmp/f038_baseline/resource_import_8e609ca.py"
)
BACKEND_DIR = Path(__file__).resolve().parents[2]
REPO_ROOT = BACKEND_DIR.parent

ROW_LIMIT_BODY = {
    "error": {
        "code": "IMPORT_ROW_LIMIT_EXCEEDED",
        "message": "Import row limit exceeded",
        "details": {"limit_rows": CAP_ROWS},
    }
}


# --- payload builders (all in memory) ----------------------------------------


def headers_for(path: str) -> list[str]:
    """每个端点使用自己的冻结模板表头（BM 14 / 26 列）。"""
    return W26_HEADERS if "resource-interfaces-ips" in path else BM_HEADERS


def _rows(headers: list[str], records: list[dict]) -> list[list[str]]:
    return [
        ["" if record.get(name) is None else str(record.get(name)) for name in headers]
        for record in records
    ]


def csv_bytes(headers: list[str], records: list[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(headers)
    for row in _rows(headers, records):
        writer.writerow(row)
    return buffer.getvalue().encode("utf-8")


def raw_csv_bytes(headers: list[str], rows: list[list[str]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return buffer.getvalue().encode("utf-8")


def xlsx_bytes(headers: list[str], rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def bm_row(index: int, *, cluster: str = CLUSTER) -> list[str]:
    return [
        cluster,
        f"f38-{index:06d}",
        "BARE_METAL",
        "",
        "",
        "Dell",
        "PowerEdge R750",
        f"SN-{index:06d}",
        "Xeon Platinum 8358",
        "2",
        "64",
        "262144",
        "A100",
        "8",
    ]


def bm_records(count: int, *, cluster: str = CLUSTER) -> list[dict]:
    out = []
    for index in range(count):
        values = bm_row(index, cluster=cluster)
        out.append(dict(zip(BM_HEADERS, values)))
    return out


def unified_row(
    index: int, *, cluster: str = CLUSTER, network: str = NETWORK, host: str | None = None
) -> list[str]:
    hostname = host or f"n{index:04d}"
    return [
        cluster,
        hostname,
        "BARE_METAL",
        "",
        "",
        "",
        "",
        "131072",
        "",
        "",
        "Dell",
        "PowerEdge R750",
        f"SN-{index:06d}",
        "Xeon Platinum 8358",
        "2",
        "64",
        "A100",
        "8",
        "eth0",
        "",
        "",
        "",
        "",
        "TRUE",
        network,
        f"10.38.0.{index + 11}",
    ]


def wide_xlsx(n_rows: int, width: int) -> bytes:
    """S-W1：`dimension` 被改写为 ``A1:<width><n+1>`` 的宽表 xlsx（仅 A 列有值）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(BM_HEADERS)
    for index in range(n_rows):
        sheet.append([f"row-{index}"] + [None] * (len(BM_HEADERS) - 1))
    buffer = io.BytesIO()
    workbook.save(buffer)
    raw = buffer.getvalue()
    return _rewrite_sheet(
        raw,
        dimension=f"A1:{get_column_letter(width)}{n_rows + 1}",
    )


def blank_range_xlsx(n_rows: int, width: int) -> bytes:
    """S-W2：表头 + ``n_rows`` 个空 ``<row/>``，`dimension` 加宽到 ``width``。

    openpyxl 会按 `dimension` 把**每一行（含表头行）**补齐到 `width` 格
    （architecture §13.4）。数据行全空 ⇒ 全属尾部空行 ⇒ `total_rows=0`，
    **不触发 413**；表头多出的空串落进 `unknown_columns` ⇒ 422。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(BM_HEADERS)
    buffer = io.BytesIO()
    workbook.save(buffer)
    raw = buffer.getvalue()
    extra = b"".join(f'<row r="{i}"/>'.encode() for i in range(2, n_rows + 2))
    return _rewrite_sheet(
        raw,
        dimension=f"A1:{get_column_letter(width)}{n_rows + 1}",
        extra_rows=extra,
    )


def _rewrite_sheet(
    raw: bytes,
    *,
    dimension: str | None = None,
    extra_rows: bytes = b"",
    replace: tuple[bytes, bytes] | None = None,
) -> bytes:
    """重写 `xl/worksheets/sheet1.xml`（`dimension` / sheetData 尾部 / 字节替换）。"""
    source = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                if extra_rows:
                    assert b"</sheetData>" in payload
                    payload = payload.replace(
                        b"</sheetData>", extra_rows + b"</sheetData>"
                    )
                if dimension is not None:
                    payload, count = re.subn(
                        rb'<dimension ref="[^"]*"/>',
                        f'<dimension ref="{dimension}"/>'.encode(),
                        payload,
                    )
                    assert count == 1, f"dimension rewrite failed: {count}"
                if replace is not None:
                    old, new = replace
                    assert old in payload, f"pattern {old!r} missing"
                    payload = payload.replace(old, new)
            target.writestr(item, payload)
    return out.getvalue()


def wide_csv(n_rows: int, width: int, *, nonblank: bool = False) -> bytes:
    """S-W6：表头 `width` 列 + `n_rows` 行。

    ``nonblank=False`` 的行全为空字段（属尾部空行，走表头/列类 422）；
    ``nonblank=True`` 的行首列非空（计入行数，超过 5000 走 413）。
    """
    header = BM_HEADERS + [""] * (width - len(BM_HEADERS))
    if nonblank:
        row = "F38" + "," * (width - 1)
    else:
        row = "," * (width - 1)
    return (",".join(header) + "\n" + (row + "\n") * n_rows).encode("utf-8")


def trailing_blank_lines(payload: bytes, blanks: int = 100_000) -> bytes:
    """S-W3：在合法 CSV 之后追加 `blanks` 个换行（纯尾部空记录）。"""
    return payload + b"\n" * blanks


def row_error_tail_csv(n_rows: int) -> bytes:
    """S-W4：`n_rows` 合法行 + 末行行级错误（非法 role）⇒ 413（不是 422）。"""
    rows = [bm_row(i) for i in range(n_rows - 1)]
    bad = bm_row(n_rows - 1)
    bad[1] = f"f38-bad-{n_rows}"
    bad[3] = "NOT_A_ROLE"
    rows.append(bad)
    return raw_csv_bytes(BM_HEADERS, rows)


def oversized_row_tail_csv(n_rows: int) -> bytes:
    """S-W5（413 的配对构造）：`n_rows` 合法行 + 末行 200 列（超宽但不解析失败）。"""
    rows = [bm_row(i) for i in range(n_rows)]
    rows.append([f"f38-wide-{n_rows}"] + [""] * 199)
    return raw_csv_bytes(BM_HEADERS, rows)


def oversized_field_tail_csv(n_rows: int, *, pad: int = 131_073) -> bytes:
    """S-W5（422 的配对构造）：`n_rows` 合法行 + 末行单字段超过 field_size_limit。"""
    rows = [bm_row(i) for i in range(n_rows)]
    rows.append(["x" * pad] + [""] * 13)
    return raw_csv_bytes(BM_HEADERS, rows)


def middle_blanks_csv(blanks: int = 1000) -> bytes:
    """S-W7：表头 + 1 合法行 + `blanks` 个空行 + 1 行非法 role。"""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(BM_HEADERS)
    writer.writerow(bm_row(0))
    for _ in range(blanks):
        writer.writerow([])
    writer.writerow(["F038", "f38-tail", "BARE_METAL", "NOT_A_ROLE"])
    return buffer.getvalue().encode("utf-8")


def header_error_csv(kind: str) -> bytes:
    """四类表头错误（`header_missing` / `missing_columns` / `unknown_columns` / `column_order`）。"""
    if kind == "header_missing":
        header = [""] * len(BM_HEADERS)
    elif kind == "missing_columns":
        header = BM_HEADERS[:-1]
    elif kind == "unknown_columns":
        header = BM_HEADERS + ["extra_column"]
    elif kind == "column_order":
        header = list(BM_HEADERS)
        header[0], header[1] = header[1], header[0]
    else:  # pragma: no cover - guard
        raise AssertionError(kind)
    return raw_csv_bytes(header, [bm_row(0)[: len(header)]])


def oversized_field_csv() -> bytes:
    """F037 输入 ①：单字段 131,073 字符 ⇒ `csv.Error` ⇒ 422 `unreadable`。"""
    row = [""] * len(BM_HEADERS)
    row[0] = "F37A"
    row[-1] = "x" * 131_073
    return raw_csv_bytes(BM_HEADERS, [row])


def corrupt_sheet_tail_xlsx() -> bytes:
    """F037 输入 ②：sheet 结束标签损坏（延迟到 `iter_rows` 才失败）。"""
    good = xlsx_bytes(BM_HEADERS, [bm_row(0)])
    return _rewrite_sheet(good, replace=(b"</worksheet>", b"</workshet>"))


def corrupt_manifest_xlsx() -> bytes:
    """F037 输入 ③：`[Content_Types].xml` 结束标签损坏。"""
    good = xlsx_bytes(BM_HEADERS, [bm_row(0)])
    source = zipfile.ZipFile(io.BytesIO(good))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                assert b"</Types>" in payload
                payload = payload.replace(b"</Types>", b"</Typse>")
            target.writestr(item, payload)
    return out.getvalue()


def missing_workbook_part_xlsx() -> bytes:
    """BE-01：manifest 可读但未声明 workbook 部件（`OSError` 路径）。"""
    content_types = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        b'<Default Extension="rels" ContentType="application/vnd.openxmlformats-'
        b'package.relationships+xml"/>'
        b'<Default Extension="xml" ContentType="application/xml"/>'
        b"</Types>"
    )
    relationships = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
        b'relationships"/>'
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        target.writestr("[Content_Types].xml", content_types)
        target.writestr("_rels/.rels", relationships)
    return out.getvalue()


def count_corrupt_xlsx(n_rows: int) -> bytes:
    """S-W5 xlsx 变体：`n_rows` 行 + sheet 结束标签损坏。"""
    good = xlsx_bytes(BM_HEADERS, [bm_row(i) for i in range(n_rows)])
    return _rewrite_sheet(good, replace=(b"</worksheet>", b"</workshet>"))


OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"not-a-real-xls"


# --- reference implementation (start_commit semantics, test-only) ------------


def _reference_xlsx_rows(data: bytes) -> list[tuple]:
    """`8e609ca` 的 `_parse_xlsx` 语义：全量物化 + 相同白名单作用域。"""
    if not data.startswith(b"PK\x03\x04"):
        raise ImportInvalidFile("unsupported_format")
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    failures = resource_import._PARSE_FAILURES + (InvalidFileException, OSError)
    try:
        from lxml.etree import LxmlError
    except ImportError:  # pragma: no cover - mirrors the implementation guard
        pass
    else:
        failures += (LxmlError,)

    workbook = None
    try:
        with resource_import._invalid_file_on_parse_failure("xlsx", len(data), failures):
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            if not workbook.worksheets:
                raise ImportInvalidFile("empty_file")
            return list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        if workbook is not None:
            resource_import._close_workbook(workbook)


def reference_records(data: bytes, fmt: str, *, max_rows: int | None) -> list[list[str]]:
    """start_commit（`8e609ca`）语义的 records（测试内参考实现，独立于被测实现）。

    与生产实现的差异**必须**保留（否则失去独立证据意义）：

    * 全量物化：`list(csv.reader(...))` / `list(worksheet.iter_rows(...))`；
    * 对**每个**单元格调用 `_cell_text`（生产实现只归一化保留记录）；
    * 不做 `keep_cells` 裁剪，表头与数据行都保持原宽度；
    * 尾部空行裁剪与行数上限均为**基线顺序**（裁剪 → 413）。

    复用的未改动共享件：`resource_import._cell_text`、`_PARSE_FAILURES`、
    `_invalid_file_on_parse_failure`、`_close_workbook`（由 S3 证明零 diff）。
    """
    if not data:
        raise ImportInvalidFile("empty_file")
    if fmt == "csv":
        try:
            decoded = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ImportInvalidFile("unreadable") from exc
        with resource_import._invalid_file_on_parse_failure("csv", len(data)):
            raw = list(csv.reader(io.StringIO(decoded, newline="")))
    else:
        raw = _reference_xlsx_rows(data)

    records = [[resource_import._cell_text(cell) for cell in row] for row in raw]
    if not records:
        raise ImportInvalidFile("empty_file")
    data_records = list(records[1:])
    while data_records and all(cell == "" for cell in data_records[-1]):
        data_records.pop()
    if max_rows is not None and len(data_records) > max_rows:
        raise ImportRowLimitExceeded(max_rows)
    return [records[0], *data_records]


# --- HTTP helpers ------------------------------------------------------------


def post(client, path: str, data: bytes, fmt: str = "csv", carrier: str = "file"):
    if fmt == "csv":
        name, mime = (
            (PASTED_FILE_NAME, PASTED_MIME)
            if carrier == "text"
            else (FILE_NAME, FILE_MIME)
        )
    else:
        name, mime = ("f038-file-carrier.xlsx", XLSX_MIME)
    return client.post(path, files={"file": (name, data, mime)}, data={"format": fmt})


def _patched_parse(data, fmt, **kwargs):
    return reference_records(data, fmt, max_rows=kwargs.get("max_rows"))


def both_paths(client, path: str, data: bytes, fmt: str = "csv", carrier: str = "file"):
    """同一 `(path, data, fmt)` 的**基线语义路径**与**生产路径**响应（TD-8.2）。

    先跑基线路径（`parse_records` 单点 monkeypatch 为参考实现），再跑真实路径；
    返回 `(fixed_resp, baseline_resp)`，由调用方断言 `(status, json())` 深等。
    """
    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(resource_import, "parse_records", _patched_parse)
        baseline = post(client, path, data, fmt, carrier)
    fixed = post(client, path, data, fmt, carrier)
    return fixed, baseline


def assert_equivalent(fixed, baseline, label: str) -> None:
    assert fixed.status_code == baseline.status_code, (
        f"{label}: status {fixed.status_code} != {baseline.status_code} "
        f"({fixed.text[:300]} | {baseline.text[:300]})"
    )
    assert fixed.json() == baseline.json(), (
        f"{label}: body differs\nfixed={fixed.text[:600]}\nbaseline={baseline.text[:600]}"
    )


def assert_row_limit(resp, label: str = "") -> None:
    assert resp.status_code == 413, f"{label}: {resp.status_code} {resp.text[:400]}"
    assert resp.json() == ROW_LIMIT_BODY, f"{label}: {resp.text[:400]}"


def assert_invalid_file(resp, reason: str, label: str = "") -> None:
    assert resp.status_code == 422, f"{label}: {resp.status_code} {resp.text[:400]}"
    error = resp.json()["error"]
    assert error["code"] == "IMPORT_INVALID_FILE", f"{label}: {error}"
    assert error["details"].get("reason") == reason, f"{label}: {error}"


def counts(session) -> dict[str, int]:
    return {
        table: int(session.execute(text(f"SELECT count(*) FROM {table}")).scalar_one())
        for table in COUNT_TABLES
    }


def assert_content_free(body: str, label: str) -> None:
    for forbidden in (
        MARKER,
        "Traceback",
        "<worksheet",
        ".py",
        "field larger than field limit",
        "F038",
    ):
        assert forbidden not in body, f"{label}: response leaks {forbidden!r}"


# --- measurement protocol (TD-6.2 / §13.2) ----------------------------------

_DATA_CACHE: dict[str, bytes] = {}
_MEASUREMENTS: dict[tuple[str, str], dict] = {}
_BASELINE_MODULE: dict[str, object] = {}
MIB = 1024 * 1024


def _baseline_module():
    """`/tmp` 的 start_commit 基线模块（只加载一次；不入库）。"""
    if "module" in _BASELINE_MODULE:
        return _BASELINE_MODULE["module"]
    module = None
    if Path(BASELINE_PATH).exists():
        spec = importlib.util.spec_from_file_location("f038_baseline_ri", BASELINE_PATH)
        assert spec is not None and spec.loader is not None
        loaded = importlib.util.module_from_spec(spec)
        # 必须注册：dataclass() 通过 sys.modules 解析本模块
        sys.modules[spec.name] = loaded
        spec.loader.exec_module(loaded)
        module = loaded
    _BASELINE_MODULE["module"] = module
    return module


def _verdict(fn):
    try:
        return ("ok", fn())
    except BaseException as exc:  # noqa: BLE001 - 仅测试内捕获，用于等价比较
        return ("exc", type(exc).__name__, getattr(exc, "details", None))


def _measure_once(fn, *, trace: bool) -> tuple[tuple, int, float]:
    gc.collect()  # 在 start() 之前
    if trace:
        tracemalloc.start()
        tracemalloc.get_traced_memory()
    started = time.perf_counter()
    verdict = _verdict(fn)
    seconds = time.perf_counter() - started
    peak = 0
    if trace:
        _current, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()
    gc.collect()  # 取峰值之后
    return verdict, peak, seconds


def _data(key: str) -> bytes:
    return _DATA_CACHE[key]


def _measure(key: str, side: str, fmt: str, *, reps: int = 2) -> dict:
    """同进程、固定顺序、两次重复；断言取两次的**较大值**（保守）。"""
    cache_key = (key, side)
    if cache_key in _MEASUREMENTS:
        return _MEASUREMENTS[cache_key]
    data = _data(key)
    if side == "baseline":
        module = _baseline_module()
        assert module is not None, "baseline module missing"
        fn = lambda: module.parse_records(data, fmt)  # noqa: E731
    else:
        fn = lambda: resource_import.parse_records(  # noqa: E731
            data, fmt, max_rows=CAP_ROWS, keep_cells=resource_import._MAX_TEMPLATE_COLUMNS
        )
    peaks: list[int] = []
    traced: list[float] = []
    raw: list[float] = []
    verdicts: list[tuple] = []
    for _ in range(reps):
        verdict, peak, seconds = _measure_once(fn, trace=True)
        peaks.append(peak)
        traced.append(seconds)
        verdicts.append(verdict)
        verdict_raw, _peak, seconds_raw = _measure_once(fn, trace=False)
        raw.append(seconds_raw)
        verdicts.append(verdict_raw)
    result = {
        "side": side,
        "fmt": fmt,
        "bytes": len(data),
        "peaks": peaks,
        "peak": max(peaks),
        "time_traced": max(traced),
        "time_traced_runs": traced,
        "time_raw": max(raw),
        "verdict": verdicts[0],
    }
    _MEASUREMENTS[cache_key] = result
    print(
        f"[M] {key:18s} {side:8s} bytes={result['bytes']:9d} "
        f"peak={result['peak'] / MIB:9.3f}MiB peaks={[round(p / MIB, 3) for p in peaks]} "
        f"t_traced={result['time_traced']:.3f}s t_raw={result['time_raw']:.3f}s "
        f"verdict={result['verdict'][0]}/{result['verdict'][1] if result['verdict'][0] == 'exc' else '-'}"
    )
    return result


_FLOOR_CACHE: dict[int, list[int]] = {}
#: 固定 `id(data)`，保证缓存键在进程内有效（样例字节由 `_DATA_CACHE` 持有）。
_FLOOR_PINNED: dict[int, bytes] = {}


def _floor_probe_once(data: bytes) -> None:
    """§14.2 `F(x)` 探针本体：`decode` → `io.StringIO` → 遍历丢弃（唯一允许的三步）。"""
    text = data.decode("utf-8-sig")
    for _record in csv.reader(io.StringIO(text, newline="")):
        pass


def _floor_probe_repeats(data: bytes) -> list[int]:
    """§14.2 地板探针两次重复峰值（M6 协议：`gc.collect()` → `start()` → 取峰值）。"""
    key = id(data)
    if key not in _FLOOR_CACHE:
        _FLOOR_PINNED[key] = data
        peaks: list[int] = []
        for _ in range(2):
            _probe_verdict, peak, _seconds = _measure_once(
                lambda: _floor_probe_once(data), trace=True
            )
            peaks.append(peak)
        _FLOOR_CACHE[key] = peaks
        print(
            f"[floor-probe] bytes={len(data)} peaks={[round(p / MIB, 3) for p in peaks]}MiB "
            f"F/B={max(peaks) / len(data):.3f}"
        )
    return _FLOOR_CACHE[key]


def _floor_probe_peak(data: bytes) -> int:
    """§14.2 `F(x)`：两次取**较大值**（保守）。"""
    return max(_floor_probe_repeats(data))


def floor_bytes(data: bytes, fmt: str) -> int:
    """§14.2 / §14.3 `F(x)`：CSV 用实测探针；xlsx 退回 1×len 的保守常量。"""
    if fmt != "csv":
        return len(data)
    return _floor_probe_peak(data)


def excess(result: dict, *, floor: int | None = None) -> int:
    """§14 口径 `work = max(peak − F, 0)`；省略 `floor` 时退回旧 1×len 口径。"""
    if floor is None:
        floor = result["bytes"]
    return max(result["peak"] - floor, 0)


def _work(key: str, result: dict) -> int:
    """按样例取地板后计算工作集：CSV = 实测 `F(x)`；xlsx = `floor=None`（1×B）。"""
    return excess(result, floor=floor_bytes(_data(key), result["fmt"]))


# =============================================================================
# U 层
# =============================================================================


@pytest.mark.parametrize(
    "values",
    [
        (),
        [],
        [],
        ("",),
        (None,),
        (" ",),
        (0,),
        (0.0,),
        (-0.0,),
        (True,),
        (False,),
        (Decimal("0"),),
        (datetime(2024, 1, 1),),
        (date(2024, 1, 1),),
        (dtime(12, 0),),
        (timedelta(0),),
        (float("nan"),),
        ("", None, ""),
        ("0",),
        (b"",),
    ],
    ids=lambda value: repr(value),
)
def test_td_f038_u1_is_blank_record_matches_cell_text(values):
    """TD-F038-U1（B3 / R-2）：与 `all(_cell_text(c) == "" ...)` 全类型等价。"""
    row = list(values)
    expected = all(resource_import._cell_text(cell) == "" for cell in row)
    assert resource_import._is_blank_record(row) is expected


def test_td_f038_u1_real_parser_products_equivalence():
    """TD-F038-U1（续）：真实解析器产出的整行空记录（CSV `[]` / xlsx 大使用区域）。"""
    empty_csv_row = next(iter(csv.reader(io.StringIO("\n"))))
    assert empty_csv_row == []
    assert resource_import._is_blank_record(empty_csv_row) is True

    wide_blank = (None,) * 16_384
    assert resource_import._is_blank_record(wide_blank) is True
    assert all(resource_import._cell_text(cell) == "" for cell in wide_blank)

    # xlsx 尾部空使用区域行 = 全 None；含 0 的行不是空行（0 归一化为 "0"）
    assert resource_import._is_blank_record((None, None, 0, None)) is False
    assert resource_import._is_blank_record(("", None, " ")) is False


def test_td_f038_u2_collect_and_finalise_semantics():
    """TD-F038-U2（B2/B3/C1–C8）：`_collect_records` + `_finalise` 语义表。"""
    collect = resource_import._collect_records
    finalise = resource_import._finalise

    # ① 行数 < cap：全部保留，返回 [header, *rows]
    rows = [["x"] * 14 for _ in range(3)]
    collected = collect(iter([["h"] * 14, *rows]), max_rows=5, keep_cells=14)
    assert collected.overflow is False
    assert len(collected.rows) == 3
    assert finalise(collected, max_rows=5) == [["h"] * 14, *rows]

    # ② 行数 == cap（max_rows+1）：不溢出；经 _finalise 才 413
    rows = [["x"] * 14 for _ in range(6)]
    collected = collect(iter([["h"] * 14, *rows]), max_rows=5, keep_cells=14)
    assert collected.overflow is False
    assert len(collected.rows) == 6
    with pytest.raises(ImportRowLimitExceeded) as excinfo:
        finalise(collected, max_rows=5)
    assert excinfo.value.details == {"limit_rows": 5}

    # ③ 行数 > cap：overflow=True 且已释放
    rows = [["x"] * 14 for _ in range(9)]
    collected = collect(iter([["h"] * 14, *rows]), max_rows=5, keep_cells=14)
    assert collected.overflow is True
    assert collected.rows == []
    with pytest.raises(ImportRowLimitExceeded):
        finalise(collected, max_rows=5)

    # ④ 中间 1 行空：以 `()` 占位补回并保持位置，_finalise 归一化为 []
    collected = collect(
        iter([["h"] * 14, ["a"] * 14, [""] * 14, ["b"] * 14]),
        max_rows=5,
        keep_cells=14,
    )
    assert collected.rows[1] == ()
    result = finalise(collected, max_rows=5)
    assert result == [["h"] * 14, ["a"] * 14, [], ["b"] * 14]

    # ⑤ 尾部 3 行空：不计入 rows，长度不位移
    collected = collect(
        iter([["h"] * 14, ["a"] * 14, [""] * 14, [""] * 14, [""] * 14]),
        max_rows=5,
        keep_cells=14,
    )
    assert len(collected.rows) == 1
    assert finalise(collected, max_rows=5) == [["h"] * 14, ["a"] * 14]

    # ⑥ 仅表头
    collected = collect(iter([["h"] * 14]), max_rows=5, keep_cells=14)
    assert finalise(collected, max_rows=5) == [["h"] * 14]

    # ⑦ header=None（输入为空记录序列）
    collected = collect(iter([]), max_rows=5, keep_cells=14)
    assert finalise(collected, max_rows=5) == []

    # ⑧ max_rows=None：无上限
    collected = collect(
        iter([["h"] * 14] + [["x"] * 14] * 20), max_rows=None, keep_cells=14
    )
    assert collected.overflow is False
    assert len(finalise(collected, max_rows=None)) == 21

    # ⑨ keep_cells=None：数据行宽度 == 输入宽度
    collected = collect(iter([["h"] * 100, ["x"] * 100]), max_rows=5, keep_cells=None)
    assert finalise(collected, max_rows=5)[1] == ["x"] * 100

    # ⑩ 100 列宽行 + keep_cells=14：数据行 14 格，表头整行（D2）
    collected = collect(iter([["h"] * 100, ["x"] * 100]), max_rows=5, keep_cells=14)
    result = finalise(collected, max_rows=5)
    assert len(result[0]) == 100
    assert len(result[1]) == 14


def test_td_f038_u3_drain_semantics_parse_error_wins(monkeypatch):
    """TD-F038-U3（B6/D4/R-7）：解析器异常不得被行数上限改写，迭代器被排空。"""
    collect = resource_import._collect_records

    consumed = {"count": 0}

    def raising():
        for index in range(20_000):
            consumed["count"] += 1
            if index == 2:
                raise csv.Error("boom")
            yield ["x"] * 14

    with pytest.raises(csv.Error):
        collect(raising(), max_rows=CAP_ROWS, keep_cells=14)
    assert consumed["count"] == 3, "异常必须在原位置抛出，且不早退"

    # `_finalise` 从未被调用（spy 计数 == 0）：走真实 `parse_records` 路径
    calls = {"count": 0}
    original = resource_import._finalise

    def spy(collected, *, max_rows):
        calls["count"] += 1
        return original(collected, max_rows=max_rows)

    monkeypatch.setattr(resource_import, "_finalise", spy)
    with pytest.raises(ImportInvalidFile) as excinfo:
        resource_import.parse_records(
            oversized_field_tail_csv(CAP_ROWS + 1),
            "csv",
            max_rows=CAP_ROWS,
            keep_cells=14,
        )
    assert excinfo.value.details == {"reason": "unreadable"}
    assert isinstance(excinfo.value.__cause__, csv.Error)
    assert calls["count"] == 0, "_finalise 不得在解析异常路径上执行"

    # 另一变体：5003 条记录（1 表头 + 5002 数据行）→ 迭代器被完整消费（计数器证明排空）
    counted = {"count": 0}

    def counted_records():
        for _ in range(CAP_ROWS + 3):
            counted["count"] += 1
            yield ["x"] * 14

    collected = collect(counted_records(), max_rows=CAP_ROWS, keep_cells=14)
    assert counted["count"] == CAP_ROWS + 3
    assert collected.overflow is True
    with pytest.raises(ImportRowLimitExceeded):
        original(collected, max_rows=CAP_ROWS)


def test_td_f038_u4_consumer_index_source_is_template_columns():
    """TD-F038-U4（静态侧）：逐行校验的下标只来自 `template.columns` / `COLUMNS`。"""
    import_source = ast.parse(
        (BACKEND_DIR / "app/services/resource_import.py").read_text()
    )
    unified_source = ast.parse(
        (BACKEND_DIR / "app/services/resource_interface_ip_import.py").read_text()
    )

    def _subscript_sources(tree, function_name):
        for node in ast.walk(tree):
            if isinstance(node, ast.FunctionDef) and node.name == function_name:
                return [
                    ast.unparse(item.slice)
                    for item in ast.walk(node)
                    if isinstance(item, ast.Subscript)
                    and isinstance(item.value, ast.Name)
                    and item.value.id in {"padded", "cells"}
                ]
        raise AssertionError(function_name)

    assert _subscript_sources(import_source, "_validate_record") == ["index"]
    assert _subscript_sources(unified_source, "_cells_for") == ["index"]

    for tree, function_name in (
        (import_source, "_validate_record"),
        (unified_source, "_cells_for"),
    ):
        node = next(
            item
            for item in ast.walk(tree)
            if isinstance(item, ast.FunctionDef) and item.name == function_name
        )
        body = ast.unparse(node)
        for token in ("padded[-1]", "cells[-1]", "max(", "[-"):
            assert token not in body, f"{function_name}: unexpected read {token}"
        assert "enumerate(template.columns)" in body or "enumerate(COLUMNS)" in body

    # `_is_blank_record` 只在 `_collect_records` 的循环内、且早于裁剪
    collect_node = next(
        node
        for node in ast.walk(import_source)
        if isinstance(node, ast.FunctionDef) and node.name == "_collect_records"
    )
    calls = [
        node
        for node in ast.walk(import_source)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Name)
        and node.func.id == "_is_blank_record"
    ]
    assert len(calls) == 1
    loop = next(
        node for node in ast.walk(collect_node) if isinstance(node, ast.For)
    )
    loop_text = ast.unparse(loop)
    assert "_is_blank_record" in loop_text
    # 裁剪出现在判定之后的语句里（文本顺序证据）
    assert loop_text.index("_is_blank_record") < loop_text.index("record[:keep_cells]")


class _PoisonTouched(AssertionError):
    """毒值被消费（architecture §13.1 P-1）。"""


class _Poison:
    """毒值：任何属性访问 / 比较 / 真假 / 文本化 / ``strip`` 都失败。

    ``len()`` **不**触发（P-1），因为 `list(cells)` 对容器取长度是合法的。
    """

    __slots__ = ()

    def _trip(self, *_args, **_kwargs):
        raise _PoisonTouched("poison cell was consumed by the consumer")

    __getattr__ = _trip
    __eq__ = _trip
    __ne__ = _trip
    __hash__ = _trip
    __bool__ = _trip
    __str__ = _trip
    __repr__ = _trip
    strip = _trip


def _poison_records(headers, rows, *, from_index, poison=None):
    records = [list(headers)]
    for row in rows:
        records.append(
            [
                (poison if poison is not None and index >= from_index else value)
                for index, value in enumerate(row)
            ]
        )
    return records


def _wide_blank_rows(rows, width):
    return [row + [""] * (width - len(row)) for row in rows]


@pytest.fixture
def poison_env(role_client_factory, auth_session_factory):
    """真实库：cluster + network，供 U4 毒值探针的成功/错误路径使用。"""
    editor, _user, _password = role_client_factory("EDITOR")
    cluster = create_cluster(editor, CLUSTER)
    create_network(editor, cluster["code"], name=NETWORK, cidr=NETWORK_CIDR)
    return {"client": editor, "session_factory": auth_session_factory}


def test_td_f038_u4_poison_probe_never_reads_beyond_keep_cells(
    poison_env, monkeypatch
):
    """TD-F038-U4（P-1..P-4）：consumer 从不读第 ≥ W 列（含错误路径 + 双负向对照）。"""
    session_factory = poison_env["session_factory"]
    settings = Settings(_env_file=None)

    bm_rows = _wide_blank_rows(
        [
            bm_row(0),
            bm_row(1, cluster="MISSING"),  # 未知 cluster
            [*bm_row(2)[:10], "x", *bm_row(2)[11:]],  # cpu_cores 非法整数
            bm_row(0),  # 文件内重复 hostname → related_row
        ],
        100,
    )
    uni_rows = _wide_blank_rows(
        [
            unified_row(0),
            unified_row(1, cluster="MISSING"),
            unified_row(2)[:25] + ["10.99.0.1"],  # 网络外地址
            unified_row(0),  # 文件内重复 (hostname, interface)
        ],
        100,
    )

    def payload(module, headers, rows, *, from_index, poison):
        records = _poison_records(headers, rows, from_index=from_index, poison=poison)
        monkeypatch.setattr(
            resource_import, "parse_records", lambda data, fmt, **kwargs: records
        )
        with session_factory() as session:
            result = module.parse_and_validate(
                session,
                b"poison",
                "csv",
                max_file_bytes=CAP_BYTES,
                max_rows=CAP_ROWS,
            )
            return resource_import.result_payload(result, settings)

    # P-2②/P-3/P-4：毒值只加数据行第 ≥ W 列 → 不触发，且与「置空」payload 完全相等
    for module, headers, rows, width in (
        (resource_import, BM_HEADERS, bm_rows, resource_import._MAX_TEMPLATE_COLUMNS),
        (unified, W26_HEADERS, uni_rows, len(unified.COLUMN_NAMES)),
    ):
        blanked = payload(module, headers, rows, from_index=width, poison=None)
        poisoned = payload(module, headers, rows, from_index=width, poison=_Poison())
        assert poisoned == blanked
        assert poisoned["error_count"] > 0, "错误路径必须被覆盖（P-4）"

        # P-2①：同一记录第 W-1 列放毒值 → 必须触发（证明探针有效、非空跑）
        with pytest.raises(_PoisonTouched):
            payload(module, headers, rows, from_index=width - 1, poison=_Poison())

    # 成功路径：全合法行第 ≥ W 列放毒值 → 不触发，importable_rows 与置空一致
    good_bm = _wide_blank_rows([bm_row(i) for i in range(3)], 100)
    blanked = payload(resource_import, BM_HEADERS, good_bm, from_index=14, poison=None)
    poisoned = payload(
        resource_import, BM_HEADERS, good_bm, from_index=14, poison=_Poison()
    )
    assert poisoned == blanked
    assert poisoned["importable_rows"] == 3
    assert poisoned["error_count"] == 0


def test_td_f038_u5_retention_upper_bound():
    """TD-F038-U5（D1/A3）：10 万条宽记录 → 排空 + 释放，保留量 ≤ max_rows+1。"""
    collected_holder = {}
    consumed = {"count": 0}

    def records():
        for _ in range(100_000):
            consumed["count"] += 1
            yield ["x"] * 200

    gc.collect()
    tracemalloc.start()
    collected = resource_import._collect_records(
        records(), max_rows=CAP_ROWS, keep_cells=14
    )
    _current, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    gc.collect()

    collected_holder["value"] = collected
    assert consumed["count"] == 100_000, "超额后仍须完整排空迭代器"
    assert collected.overflow is True
    assert collected.rows == []
    assert len(collected.rows) <= CAP_ROWS + 1
    assert peak <= 8 * MIB, f"纯生成器用例峰值 {peak / MIB:.3f}MiB 超过 8MiB 哨兵"
    with pytest.raises(ImportRowLimitExceeded) as excinfo:
        resource_import._finalise(collected, max_rows=CAP_ROWS)
    assert excinfo.value.details == {"limit_rows": CAP_ROWS}


def test_td_f038_u6_two_argument_calls_stay_compatible():
    """TD-F038-U6（DR-F038-02 / ON-2 / E-3 / B12）：两参数调用与显式 `None` 等价。"""
    import inspect

    signature = inspect.signature(resource_import.parse_records)
    assert list(signature.parameters) == ["data", "fmt", "max_rows", "keep_cells"]
    for name in ("max_rows", "keep_cells"):
        assert signature.parameters[name].default is None
        assert signature.parameters[name].kind is inspect.Parameter.KEYWORD_ONLY

    samples = {
        "legal_csv": (
            csv_bytes(BM_HEADERS, bm_records(3)),
            "csv",
            BM_HEADERS,
        ),
        "legal_xlsx": (
            xlsx_bytes(BM_HEADERS, [bm_row(i) for i in range(3)]),
            "xlsx",
            BM_HEADERS,
        ),
        "trailing_blanks": (
            trailing_blank_lines(csv_bytes(BM_HEADERS, bm_records(2)), 5),
            "csv",
            BM_HEADERS,
        ),
        "middle_blank": (
            raw_csv_bytes(BM_HEADERS, [bm_row(0), [""] * 14, bm_row(1)]),
            "csv",
            BM_HEADERS,
        ),
        "wide_header": (
            wide_csv(1, 100),
            "csv",
            BM_HEADERS + [""] * (100 - len(BM_HEADERS)),
        ),
    }
    for label, (payload, fmt, expected_header) in samples.items():
        two_args = resource_import.parse_records(payload, fmt)
        explicit = resource_import.parse_records(
            payload, fmt, max_rows=None, keep_cells=None
        )
        assert isinstance(two_args, list)
        assert two_args == explicit, label
        assert two_args[0] == expected_header, f"{label}: 表头必须完整保留"
        assert all(isinstance(cell, str) for row in two_args for cell in row), label


# =============================================================================
# I 层（真实 PostgreSQL + 真实 create_app）
# =============================================================================


@pytest.fixture
def editor_client(role_client_factory):
    editor, _user, _password = role_client_factory("EDITOR")
    return editor


@pytest.fixture
def cluster_env(editor_client):
    """BM14 端点用的既有 Cluster。"""
    return create_cluster(editor_client, CLUSTER)


@pytest.fixture
def unified_env(editor_client, cluster_env):
    """26 列端点用的既有 Cluster + Network（复用 `cluster_env` 的 Cluster）。"""
    create_network(editor_client, cluster_env["code"], name=NETWORK, cidr=NETWORK_CIDR)
    return cluster_env


def test_td_f038_i1_wide_sheet_row_limit_every_endpoint(editor_client):
    """TD-F038-I1（S-W1 / A1–A3 / B2）：702 列 × 5001 行 xlsx → 4 端点 413 且差分相等。"""
    payload = wide_xlsx(CAP_ROWS + 1, 702)
    assert payload.startswith(b"PK\x03\x04")
    assert len(payload) < CAP_BYTES
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        assert workbook.worksheets[0].max_column == 702
        assert workbook.worksheets[0].max_row == CAP_ROWS + 2
    finally:
        workbook.close()

    for path in ALL_PATHS:
        fixed, baseline = both_paths(editor_client, path, payload, "xlsx")
        assert_equivalent(fixed, baseline, f"{path} S-W1")
        assert_row_limit(fixed, f"{path} S-W1")


@pytest.mark.parametrize("rows", [101], ids=["rows101"])
def test_td_f038_i2_large_used_range_is_unknown_columns(editor_client, rows):
    """TD-F038-I2（S-W2 / §13.4 TD-I-8 勘误）：16,384 列空使用区域 → **422 unknown_columns**。

    期望值依据：`dimension A1:XFD{N+1}` 使 openpyxl 把表头行补齐到 16,384 格
    （14 个真实列名 + 16,370 个空串）；数据行全空 ⇒ 全部作为尾部空行裁掉 ⇒
    `total_rows=0` ⇒ **不触发 413**。规模在这里缩小到 `rows` 行（响应与行数无关，
    已在本用例断言），原规模测量见 M4。
    """
    payload = blank_range_xlsx(rows, 16_384)
    assert payload.startswith(b"PK\x03\x04")
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        assert worksheet.max_column == 16_384
        first = next(worksheet.iter_rows(values_only=True))
        assert len(first) == 16_384
        assert list(first[:14]) == BM_HEADERS
        assert set(first[14:]) == {None}
    finally:
        workbook.close()

    for path in ALL_PATHS:
        fixed, baseline = both_paths(editor_client, path, payload, "xlsx")
        assert_equivalent(fixed, baseline, f"{path} S-W2")
        assert fixed.status_code == 422, f"{path}: {fixed.status_code} {fixed.text[:300]}"
        assert fixed.json()["error"]["details"]["reason"] == "unknown_columns"
        assert fixed.json()["error"]["code"] != "IMPORT_ROW_LIMIT_EXCEEDED"
        unknown = fixed.json()["error"]["details"].get("unknown_columns")
        assert unknown is not None and set(unknown) == {""}
        assert len(unknown) == 16_384 - len(BM_HEADERS)


def test_td_f038_i3_trailing_blank_lines_do_not_move_row_numbers(
    editor_client, unified_env, auth_session_factory
):
    """TD-F038-I3（S-W3 / A4 / B3 / T113）：尾部 10 万个空行不计入 `total_rows`。"""
    cases = (
        (RES_VALIDATE, RES_WRITE, csv_bytes(BM_HEADERS, bm_records(3))),
        (
            UNI_VALIDATE,
            UNI_WRITE,
            raw_csv_bytes(W26_HEADERS, [unified_row(i) for i in range(3)]),
        ),
    )
    for validate_path, write_path, body in cases:
        payload = trailing_blank_lines(body, 100_000)
        assert len(payload) < CAP_BYTES
        fixed, baseline = both_paths(editor_client, validate_path, payload, "csv")
        assert_equivalent(fixed, baseline, validate_path)
        assert fixed.status_code == 200, fixed.text[:300]
        body_json = fixed.json()
        assert body_json["total_rows"] == 3
        assert body_json["error_count"] == 0
        assert body_json["importable_rows"] == 3
        assert body_json["errors"] == []
        assert body_json["errors_truncated"] is False

        response = post(editor_client, write_path, payload, "csv")
        assert response.status_code == 201, response.text[:300]
        assert response.json()["total_rows"] == 3
        assert response.json()["created"] == 3


@pytest.mark.parametrize("variant", ["row_error", "wide_row"])
def test_td_f038_i4_row_level_tail_error_stays_413(editor_client, variant):
    """TD-F038-I4（S-W4 / ON-1 / B6）：尾部**行级错误/超宽行** → 4 端点 413。"""
    payload = (
        row_error_tail_csv(CAP_ROWS + 1)
        if variant == "row_error"
        else oversized_row_tail_csv(CAP_ROWS + 1)
    )
    assert len(payload) < CAP_BYTES
    for path in ALL_PATHS:
        fixed, baseline = both_paths(editor_client, path, payload, "csv")
        assert_equivalent(fixed, baseline, f"{path} S-W4 {variant}")
        assert_row_limit(fixed, f"{path} S-W4 {variant}")


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_td_f038_i5_tail_parser_error_is_422_unreadable(editor_client, fmt):
    """TD-F038-I5（S-W5 / ON-1 / B8）：尾部**解析器异常** → 4 端点 422 `unreadable`。"""
    payload = (
        oversized_field_tail_csv(CAP_ROWS + 1)
        if fmt == "csv"
        else count_corrupt_xlsx(CAP_ROWS + 1)
    )
    assert len(payload) < CAP_BYTES
    for path in ALL_PATHS:
        fixed, baseline = both_paths(editor_client, path, payload, fmt)
        assert_equivalent(fixed, baseline, f"{path} S-W5 {fmt}")
        assert fixed.status_code == 422, f"{path}: {fixed.status_code} {fixed.text[:300]}"
        assert fixed.json()["error"]["details"] == {"reason": "unreadable"}
        assert fixed.json()["error"]["code"] != "IMPORT_ROW_LIMIT_EXCEEDED"


def test_td_f038_i6_wide_csv_stays_a_column_reason(editor_client):
    """TD-F038-I6（S-W6 / A7 / B7 / N1）：宽表 CSV 走既有列类 reason，不新增拒绝。"""
    wide_blank = wide_csv(4000, 651)
    assert len(wide_blank) < CAP_BYTES
    for path in ALL_PATHS:
        fixed, baseline = both_paths(editor_client, path, wide_blank, "csv")
        assert_equivalent(fixed, baseline, f"{path} S-W6a")
        assert fixed.status_code == 422, f"{path}: {fixed.status_code} {fixed.text[:300]}"
        details = fixed.json()["error"]["details"]
        assert details["reason"] == "unknown_columns"
        assert len(details["unknown_columns"]) == 651 - len(BM_HEADERS)
        assert set(details["unknown_columns"]) == {""}

    # b：5001 行**非空**宽表（行数超限）→ 413
    wide_rows = wide_csv(CAP_ROWS + 1, 651, nonblank=True)
    assert len(wide_rows) < CAP_BYTES
    for path in ALL_PATHS:
        fixed, baseline = both_paths(editor_client, path, wide_rows, "csv")
        assert_equivalent(fixed, baseline, f"{path} S-W6b")
        assert_row_limit(fixed, f"{path} S-W6b")


def test_td_f038_i7_middle_blanks_keep_row_numbers(
    editor_client, cluster_env, auth_session_factory
):
    """TD-F038-I7（S-W7 / A4 / R-1 / R-4 / B5）：1000 个中间空行不移动行号。"""
    payload = middle_blanks_csv(1000)
    with auth_session_factory() as session:
        before = counts(session)

    fixed, baseline = both_paths(editor_client, RES_VALIDATE, payload, "csv")
    assert_equivalent(fixed, baseline, RES_VALIDATE)
    assert fixed.status_code == 200, fixed.text[:300]
    body = fixed.json()
    assert body["total_rows"] == 1002
    assert body["error_count"] == 1000 * 3 + 1
    assert body["errors_truncated"] is True
    assert len(body["errors"]) == CAP_PREVIEW
    assert body["importable_rows"] == 1
    assert body["errors"][0]["row"] == 3, "首个中间空行的行号必须是 3"
    assert {error["row"] for error in body["errors"]} <= set(range(3, 1004))

    # 小规模配对（不触发预览截断）：末行行号必须仍是 6（1000 空行不位移）
    small = middle_blanks_csv(3)
    small_fixed, small_baseline = both_paths(editor_client, RES_VALIDATE, small, "csv")
    assert_equivalent(small_fixed, small_baseline, "S-W7 small")
    small_body = small_fixed.json()
    assert small_body["total_rows"] == 5
    assert small_body["error_count"] == 3 * 3 + 1
    assert small_body["errors_truncated"] is False
    assert small_body["importable_rows"] == 1
    assert {error["row"] for error in small_body["errors"]} == {3, 4, 5, 6}

    write_resp = post(editor_client, RES_WRITE, payload, "csv")
    assert write_resp.status_code == 422, write_resp.text[:300]
    assert write_resp.json()["error"]["code"] == "IMPORT_VALIDATION_FAILED"
    with auth_session_factory() as session:
        assert counts(session) == before


def test_td_f038_i8_legal_max_rows_validate(editor_client, cluster_env):
    """TD-F038-I8（A5 / B1 / B4）：合法 5000 行 → 200 且计数口径不变。"""
    payload = csv_bytes(BM_HEADERS, bm_records(CAP_ROWS))
    assert len(payload) < CAP_BYTES
    fixed, baseline = both_paths(editor_client, RES_VALIDATE, payload, "csv")
    assert_equivalent(fixed, baseline, "legal5000")
    assert fixed.status_code == 200, fixed.text[:300]
    body = fixed.json()
    assert body["total_rows"] == CAP_ROWS
    assert body["error_count"] == 0
    assert body["importable_rows"] == CAP_ROWS
    assert body["errors"] == []
    assert body["errors_truncated"] is False


def test_td_f038_i9_rejected_paths_leave_no_residue(
    editor_client, cluster_env, unified_env, auth_session_factory
):
    """TD-F038-I9（A9 / B10 / B13）：拒绝路径 7 表零变化 + 响应内容保护。"""
    samples = {
        "csv": (
            row_error_tail_csv(CAP_ROWS + 1),
            oversized_field_tail_csv(CAP_ROWS + 1),
            wide_csv(4000, 651),
        ),
        "xlsx": (
            wide_xlsx(CAP_ROWS + 1, 702),
            blank_range_xlsx(101, 16_384),
            count_corrupt_xlsx(3),
        ),
    }
    with auth_session_factory() as session:
        before = counts(session)

    responses = []
    for fmt, payloads in samples.items():
        for payload in payloads:
            for path in ALL_PATHS:
                response = post(editor_client, path, payload, fmt)
                responses.append((path, fmt, response))

    for path, fmt, response in responses:
        assert response.status_code in (413, 422), (
            f"{path} {fmt}: {response.status_code} {response.text[:200]}"
        )
        assert_content_free(response.text, f"{path} {fmt}")

    with auth_session_factory() as session:
        after = counts(session)
    assert after == before, f"拒绝路径不得有业务副作用: {before} -> {after}"


def test_td_f038_i10_unauthenticated_precedes_bounded_parse(auth_app):
    """TD-F038-I10（A10 / B11 / T124）：401 先于有界解析（不消费请求体改变顺序）。"""
    payloads = (
        wide_xlsx(CAP_ROWS + 1, 702),
        blank_range_xlsx(101, 16_384),
    )

    unauth = TestClient(auth_app)
    unauth.__enter__()
    try:
        for payload in payloads:
            for path in ALL_PATHS:
                response = post(unauth, path, payload, "xlsx")
                assert response.status_code == 401, (
                    f"{path}: {response.status_code} {response.text[:200]}"
                )
                assert response.json()["error"]["code"] == "AUTH_REQUIRED"
    finally:
        unauth.__exit__(None, None, None)


def test_td_f038_i10b_viewer_is_forbidden(role_client_factory):
    """TD-F038-I10（续）：VIEWER 403 先于有界解析，不出现 413/422/500。"""
    viewer, _user, _password = role_client_factory("VIEWER")
    payloads = (
        wide_xlsx(CAP_ROWS + 1, 702),
        blank_range_xlsx(101, 16_384),
    )
    for payload in payloads:
        for path in ALL_PATHS:
            response = post(viewer, path, payload, "xlsx")
            assert response.status_code == 403, (
                f"{path}: {response.status_code} {response.text[:200]}"
            )
            assert response.json()["error"] == {
                "code": "FORBIDDEN",
                "message": "Forbidden",
                "details": {"reason": "permission"},
            }


def test_td_f038_i11_limit_diagonal_order_unchanged(editor_client, monkeypatch):
    """TD-F038-I11（A5 / A6 / B1 / B2 / B6）：413 字节 vs 413 行数 vs 表头 422 的顺序。"""
    too_large = b"x" * (CAP_BYTES + 1)
    for path in ALL_PATHS:
        response = post(editor_client, path, too_large, "csv")
        assert response.status_code == 413, f"{path}: {response.status_code}"
        assert response.json() == {
            "error": {
                "code": "IMPORT_FILE_TOO_LARGE",
                "message": "Import file too large",
                "details": {"limit_bytes": CAP_BYTES},
            }
        }

    # (b) 5001 合法行 → 413 行数（差分相等）
    legal = csv_bytes(BM_HEADERS, bm_records(CAP_ROWS + 1))
    for path in ALL_PATHS:
        fixed, baseline = both_paths(editor_client, path, legal, "csv")
        assert_equivalent(fixed, baseline, f"{path} 5001 rows")
        assert_row_limit(fixed, f"{path} 5001 rows")

    # (c) 5001 行 + 错误表头 → 行数先于表头
    bad_header = raw_csv_bytes(
        BM_HEADERS[:-1], [bm_row(i)[:-1] for i in range(CAP_ROWS + 1)]
    )
    for path in ALL_PATHS:
        fixed, baseline = both_paths(editor_client, path, bad_header, "csv")
        assert_equivalent(fixed, baseline, f"{path} 5001 + bad header")
        assert_row_limit(fixed, f"{path} 5001 + bad header")

    # (d) import_max_rows=1 + 2 行 + 错误表头 → 仍是 413（既有顺序口径）
    settings = Settings(_env_file=None, import_max_rows=1)
    two_rows_bad_header = raw_csv_bytes(
        BM_HEADERS[:-1], [bm_row(i)[:-1] for i in range(2)]
    )
    with monkeypatch.context() as patch:
        patch.setattr(imports_api, "get_settings", lambda: settings)
        for path in ALL_PATHS:
            response = post(editor_client, path, two_rows_bad_header, "csv")
            assert response.status_code == 413, f"{path}: {response.status_code}"
            assert response.json()["error"]["details"] == {"limit_rows": 1}


def test_td_f038_i12_error_preview_and_importable_rows(
    editor_client, cluster_env, auth_session_factory
):
    """TD-F038-I12（A5 / B4 / R-4 / T109）：1100 行错误 → 预览 1000 + 计数口径不变。"""
    rows = [bm_row(i, cluster="MISSING") for i in range(1100)]
    payload = raw_csv_bytes(BM_HEADERS, rows)
    with auth_session_factory() as session:
        before = counts(session)

    fixed, baseline = both_paths(editor_client, RES_VALIDATE, payload, "csv")
    assert_equivalent(fixed, baseline, "1100 errors")
    body = fixed.json()
    assert body["total_rows"] == 1100
    assert body["error_count"] == 1100
    assert len(body["errors"]) == CAP_PREVIEW
    assert body["errors_truncated"] is True
    assert body["importable_rows"] == 0
    for error in body["errors"]:
        assert set(error) == {"row", "field", "reason", "related_row"}
        assert error["field"] == "cluster"
    assert [error["row"] for error in body["errors"]] == list(range(2, 1002))

    write_resp = post(editor_client, RES_WRITE, payload, "csv")
    assert write_resp.status_code == 422, write_resp.text[:300]
    with auth_session_factory() as session:
        assert counts(session) == before


def test_td_f038_i13_dual_carrier_equivalence(editor_client, cluster_env):
    """TD-F038-I13（A8 / B9 / T107 / T114）：两载体响应逐字段相同。"""
    # 成功批次（S-W3）只比两次只读 validate：写入端点重复提交会因唯一性冲突改变
    # 第二次响应（单侧证据按 R-14 登记，写入成功路径的载体等价由 F034 R3 覆盖）。
    payloads = (
        (trailing_blank_lines(csv_bytes(BM_HEADERS, bm_records(2)), 1000), "S-W3", (RES_VALIDATE, UNI_VALIDATE)),
        (wide_csv(400, 651), "S-W6a", ALL_PATHS),
    )
    for payload, label, paths in payloads:
        for path in paths:
            file_resp = post(editor_client, path, payload, "csv", carrier="file")
            text_resp = post(editor_client, path, payload, "csv", carrier="text")
            assert file_resp.status_code == text_resp.status_code, f"{path} {label}"
            assert file_resp.json() == text_resp.json(), f"{path} {label}"
            again = post(editor_client, path, payload, "csv", carrier="text")
            assert again.json() == text_resp.json(), f"{path} {label} 重复提交"


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_td_f038_i15_f037_inputs_still_422_unreadable(editor_client, fmt):
    """TD-F038-I15（A1 / A7 / B6–B9 / T72 / T112）：F037 输入与既有 reason 全不变。"""
    if fmt == "csv":
        payloads = {
            "oversized_field": oversized_field_csv(),
            "empty": b"",
            "header_only": raw_csv_bytes(BM_HEADERS, []),
            **{f"header_{kind}": header_error_csv(kind) for kind in
               ("header_missing", "missing_columns", "unknown_columns", "column_order")},
        }
    else:
        payloads = {
            "oversized_field": oversized_field_csv(),  # 以 xlsx 提交 ⇒ unsupported_format
            "corrupt_sheet_tail": corrupt_sheet_tail_xlsx(),
            "corrupt_manifest": corrupt_manifest_xlsx(),
            "missing_workbook_part": missing_workbook_part_xlsx(),
            "ole2": OLE2_MAGIC,
            "empty": b"",
        }

    for label, payload in payloads.items():
        for path in ALL_PATHS:
            fixed, baseline = both_paths(editor_client, path, payload, fmt)
            assert_equivalent(fixed, baseline, f"{path} {label}")
            assert fixed.status_code in (422,), (
                f"{path} {label}: {fixed.status_code} {fixed.text[:200]}"
            )
            assert fixed.status_code != 500


@pytest.mark.parametrize(
    "label, payload_factory, reason",
    [
        ("oversized_field_csv", oversized_field_csv, "unreadable"),
        ("corrupt_sheet_tail", corrupt_sheet_tail_xlsx, "unreadable"),
        ("corrupt_manifest", corrupt_manifest_xlsx, "unreadable"),
        ("missing_workbook_part", missing_workbook_part_xlsx, "unreadable"),
    ],
)
def test_td_f038_i15b_unreadable_details_are_exact(
    editor_client, label, payload_factory, reason
):
    """TD-F038-I15（续 / B8）：`unreadable` 的 `details` 恰 `{"reason": "unreadable"}`。"""
    payload = payload_factory()
    fmt = "csv" if label == "oversized_field_csv" else "xlsx"
    for path in ALL_PATHS:
        fixed, baseline = both_paths(editor_client, path, payload, fmt)
        assert_equivalent(fixed, baseline, f"{path} {label}")
        assert fixed.status_code == 422, f"{path}: {fixed.status_code}"
        assert fixed.json()["error"]["details"] == {"reason": reason}


# =============================================================================
# M 层（测量：公式原样，不放宽）
# =============================================================================


def test_td_f038_m1_environment_record():
    """TD-F038-M1（A12 / ON-4 / ON-6）：环境必录项（写入 test-report §2）。"""
    import lxml
    import openpyxl

    settings = Settings(_env_file=None)
    assert settings.import_max_file_bytes == CAP_BYTES
    assert settings.import_max_rows == CAP_ROWS
    assert settings.import_error_preview_limit == CAP_PREVIEW
    assert resource_import._MAX_TEMPLATE_COLUMNS == len(BM_HEADERS)
    info = {
        "python": sys.version.split()[0],
        "platform": sys.platform,
        "openpyxl": openpyxl.__version__,
        "openpyxl.LXML": openpyxl.LXML,
        "lxml": lxml.__version__,
        "csv.field_size_limit": csv.field_size_limit(),
        "pytest": pytest.__version__,
        "head": subprocess.run(
            ["git", "-C", str(REPO_ROOT), "rev-parse", "--short", "HEAD"],
            capture_output=True,
            text=True,
            check=True,
        ).stdout.strip(),
        "baseline_module": BASELINE_PATH,
        "baseline_module_present": Path(BASELINE_PATH).exists(),
    }
    print(f"[M1] environment={info}")
    assert info["baseline_module_present"], (
        f"基线模块缺失（{BASELINE_PATH}）：M2–M6 无法执行"
    )


def test_td_f038_m2_reference_matches_start_commit_baseline():
    """TD-F038-M2（A12 / TD-I-4）：参考实现与 `/tmp` 真实基线模块黄金对照。"""
    module = _baseline_module()
    assert module is not None

    # 独立性自证（代码体，不含 docstring）：必须全量物化，且不含 F038 的新符号
    module_tree = ast.parse(Path(__file__).read_text())
    code_bodies = []
    for node in ast.walk(module_tree):
        if isinstance(node, ast.FunctionDef) and node.name in {
            "reference_records",
            "_reference_xlsx_rows",
        }:
            body = node.body
            if (
                body
                and isinstance(body[0], ast.Expr)
                and isinstance(body[0].value, ast.Constant)
                and isinstance(body[0].value.value, str)
            ):
                body = body[1:]
            code_bodies.append(ast.unparse(ast.Module(body=body, type_ignores=[])))
    reference_code = "\n".join(code_bodies)
    assert "list(csv.reader(" in reference_code
    assert "list(workbook.worksheets[0].iter_rows(" in reference_code
    for forbidden in ("_collect_records", "_finalise", "keep_cells"):
        assert forbidden not in reference_code, forbidden

    samples = {
        "legal_csv": (csv_bytes(BM_HEADERS, bm_records(3)), "csv"),
        "wide_header_csv": (wide_csv(2, 100), "csv"),
        "trailing_blanks": (trailing_blank_lines(csv_bytes(BM_HEADERS, bm_records(2)), 50), "csv"),
        "oversized_field": (oversized_field_csv(), "csv"),
        "legal_xlsx": (xlsx_bytes(BM_HEADERS, [bm_row(i) for i in range(3)]), "xlsx"),
        "wide_xlsx": (wide_xlsx(4, 702), "xlsx"),
        "corrupt_sheet_tail": (count_corrupt_xlsx(3), "xlsx"),
    }
    for label, (payload, fmt) in samples.items():
        for max_rows in (CAP_ROWS, None):
            mine = _verdict(lambda: reference_records(payload, fmt, max_rows=max_rows))
            golden = _verdict(
                lambda: _golden_records(module, payload, fmt, max_rows=max_rows)
            )
            assert mine[0] == golden[0], f"{label}: {mine[0]} != {golden[0]}"
            if mine[0] == "ok":
                assert mine[1] == golden[1], label
            else:
                assert mine[1:] == golden[1:], f"{label}: {mine[1:]} != {golden[1:]}"


def _golden_records(module, data: bytes, fmt: str, *, max_rows):
    """`/tmp` 基线模块 + 基线 `parse_and_validate` 的裁剪/行数顺序（黄金对照）。"""
    records = module.parse_records(data, fmt)
    if not records:
        raise ImportInvalidFile("empty_file")
    data_records = list(records[1:])
    while data_records and all(cell == "" for cell in data_records[-1]):
        data_records.pop()
    if max_rows is not None and len(data_records) > max_rows:
        raise ImportRowLimitExceeded(max_rows)
    return [records[0], *data_records]


_SAMPLES = {
    # key: (fmt, builder)
    "S-W1-1001": ("xlsx", lambda: wide_xlsx(1001, 702)),
    "S-W2-1001": ("xlsx", lambda: blank_range_xlsx(1001, 16_384)),
    "S-W6a-4000": ("csv", lambda: wide_csv(4000, 651)),
    "legal-5000": ("csv", lambda: csv_bytes(BM_HEADERS, bm_records(CAP_ROWS))),
    "S-W1-N2501": ("xlsx", lambda: wide_xlsx(2501, 702)),
    "S-W1-4N10001": ("xlsx", lambda: wide_xlsx(10001, 702)),
    "S-W2-N2501": ("xlsx", lambda: blank_range_xlsx(2501, 16_384)),
    "S-W2-4N10001": ("xlsx", lambda: blank_range_xlsx(10001, 16_384)),
    "S-W6a-1000": ("csv", lambda: wide_csv(1000, 651)),
    "S-W6b": ("csv", lambda: wide_csv(CAP_ROWS + 1, 651, nonblank=True)),
    "S-W6c": ("csv", lambda: wide_csv(CAP_ROWS + 1, 14, nonblank=True)),
    "legal-1000": ("csv", lambda: csv_bytes(BM_HEADERS, bm_records(1000))),
    "legal-4000": ("csv", lambda: csv_bytes(BM_HEADERS, bm_records(4000))),
}


def _sample(key: str) -> dict:
    """按 `_SAMPLES` 构造样例（构造发生在 `tracemalloc.start()` **之前**）。"""
    if key not in _DATA_CACHE:
        _DATA_CACHE[key] = _SAMPLES[key][1]()
    return _measure(key, "fixed", _SAMPLES[key][0])


def _pair(key: str) -> tuple[dict, dict]:
    fmt = _SAMPLES[key][0]
    if key not in _DATA_CACHE:
        _DATA_CACHE[key] = _SAMPLES[key][1]()
    return _measure(key, "baseline", fmt), _measure(key, "fixed", fmt)


@pytest.mark.parametrize("key", ["S-W1-1001", "S-W2-1001", "S-W6a-4000", "S-W6b"])
def test_td_f038_m3a_t2_relative_to_baseline(key):
    """TD-F038-M3（T-2′，§14.2）：`work(fixed) ≤ work(baseline) / 8`（等规模）。

    CSV 族地板取**实测探针** `floor_bytes(...)`（`F(x)`）；xlsx 族保持 `floor=None`
    （1×len 保守常量）⇒ xlsx 的数值与判定与 §13.2 逐位一致（须复现 11.44× / 155.73×）。
    G-B（§14.3）：`work(baseline) ≥ 1 MiB`，防比值退化为 `0 ≤ 0`。
    """
    baseline, fixed = _pair(key)
    floor = floor_bytes(_data(key), fixed["fmt"])
    fixed_work = excess(fixed, floor=floor)
    baseline_work = excess(baseline, floor=floor)
    print(
        f"[T-2'] {key}: bytes={fixed['bytes']} F={floor / MIB:.3f}MiB "
        f"peak_fixed={fixed['peak'] / MIB:.3f}MiB peak_baseline={baseline['peak'] / MIB:.3f}MiB "
        f"work_fixed={fixed_work / MIB:.3f}MiB work_baseline={baseline_work / MIB:.3f}MiB "
        f"bound={baseline_work / 8 / MIB:.3f}MiB "
        f"sources={(baseline_work / fixed_work) if fixed_work else float('inf'):.4f}x"
    )
    assert baseline_work >= MIB, (
        f"G-B {key}: work(baseline)={baseline_work / MIB:.3f}MiB < 1MiB（比值判据退化）"
    )
    assert fixed_work * 8 <= baseline_work, (
        f"T-2' {key}: work(fixed)={fixed_work / MIB:.3f}MiB > "
        f"work(baseline)/8={baseline_work / 8 / MIB:.3f}MiB "
        f"(实测倍数 {(baseline_work / fixed_work) if fixed_work else float('inf'):.3f}x)"
    )


def test_td_f038_m3b_t3_no_regression_on_legal_max_rows():
    """TD-F038-M3（T-3，§13.2）：合法 5000 行 `fixed_peak ≤ 1.2 × baseline_peak`。"""
    baseline, fixed = _pair("legal-5000")
    print(
        f"[T-3] legal-5000: peak_fixed={fixed['peak'] / MIB:.3f}MiB "
        f"peak_baseline={baseline['peak'] / MIB:.3f}MiB ratio={fixed['peak'] / baseline['peak']:.4f}"
    )
    assert fixed["peak"] <= 1.2 * baseline["peak"]


@pytest.mark.parametrize(
    "key", ["S-W1-1001", "S-W2-1001", "S-W6a-4000", "legal-5000"]
)
def test_td_f038_m3c_t4_absolute_sentinel(key):
    """TD-F038-M3（T-4，§13.2）：`fixed_peak ≤ 64 × len(data) + 64 MiB`。"""
    fixed = _sample(key)
    bound = 64 * fixed["bytes"] + 64 * MIB
    print(
        f"[T-4] {key}: peak_fixed={fixed['peak'] / MIB:.3f}MiB "
        f"bytes={fixed['bytes']} bound={bound / MIB:.3f}MiB"
    )
    assert fixed["peak"] <= bound


@pytest.mark.parametrize("family", ["S-W1", "S-W2"])
def test_td_f038_m4a_t1_scale_insensitivity_xlsx(family):
    """TD-F038-M4（T-1，§13.2）：xlsx 族 `fixed(4N) ≤ 1.5 × fixed(N) + 2 MiB`。"""
    small = _sample(f"{family}-N2501")
    large = _sample(f"{family}-4N10001")
    bound = 1.5 * small["peak"] + 2 * MIB
    print(
        f"[T-1] {family}: N=2501 peak={small['peak'] / MIB:.3f}MiB "
        f"(bytes={small['bytes']}, verdict={small['verdict'][0]}) "
        f"4N=10001 peak={large['peak'] / MIB:.3f}MiB "
        f"(bytes={large['bytes']}, verdict={large['verdict'][0]}) "
        f"bound={bound / MIB:.3f}MiB"
    )
    assert large["peak"] <= bound


def test_td_f038_m4b_t1b_csv_delta():
    """TD-F038-M4（T-1b′，§14.2）：CSV 族 `work(4N) ≤ 1.5 × work(N) + 2 MiB`。

    `work(x) = max(peak − F(x), 0)`，`F(x)` 为本进程同字节实测地板。旧 T-1b 的
    `Δbytes` 配额已删除（收紧）；合法 ≤5000 行族不再参战（改由 T-3 + U5 承担）。
    """
    small = _sample("S-W6a-1000")
    large = _sample("S-W6a-4000")
    work_small = _work("S-W6a-1000", small)
    work_large = _work("S-W6a-4000", large)
    bound = 1.5 * work_small + 2 * MIB
    print(
        f"[T-1b'] S-W6a: peak_1000={small['peak'] / MIB:.3f}MiB "
        f"peak_4000={large['peak'] / MIB:.3f}MiB "
        f"work_1000={work_small / MIB:.3f}MiB work_4000={work_large / MIB:.3f}MiB "
        f"bound={bound / MIB:.3f}MiB margin={(bound / work_large) if work_large else float('inf'):.1f}x"
    )
    assert work_large <= bound, (
        f"T-1b' S-W6a: work(4000)={work_large / MIB:.3f}MiB > "
        f"1.5×work(1000)+2MiB={bound / MIB:.3f}MiB"
    )


def test_td_f038_m4d_csv_peak_floor_coefficient_report():
    """TD-F038-M4（§14.2 地板来源 + §14.3 G-A1/G-A2）：产出 `F(x)` 并公布 `F/B`。

    `F(x)` 由 `_floor_probe_peak` 实测（`decode("utf-8-sig")` → `io.StringIO(text,
    newline="")` → 遍历丢弃；M6 协议、两次取较大值），**同进程同字节标定** ⇒ 公式内
    不含解释器相关常量。期望 `F/B ≈ 5.0–5.02`：CSV 路径必然经过
    `data.decode("utf-8-sig")`（1×len）+ `io.StringIO(text)`（CPython 内部 UCS-4
    缓冲 4×len），两行均为 start_commit 基线原样代码，AQ-2=A 下不可消除。

    守卫（缺一即口径不成立）：G-A1 `F ≥ B`（防探针空跑）；
    G-A2 `F ≤ 8B + 4 MiB`（防探针自身保留导致判据稀释）。
    """
    for key in ("S-W6a-1000", "S-W6a-4000", "S-W6b", "S-W6c"):
        fixed = _sample(key)
        payload = _data(key)
        floor = floor_bytes(payload, fixed["fmt"])
        repeats = _floor_probe_repeats(payload)
        work = excess(fixed, floor=floor)
        print(
            f"[floor-F] {key:12s} B={len(payload):9d} F={floor / MIB:9.3f}MiB "
            f"F/B={floor / len(payload):.3f} repeats={[round(p / MIB, 3) for p in repeats]} "
            f"peak={fixed['peak'] / MIB:.3f}MiB work={work / MIB:.3f}MiB "
            f"verdict={fixed['verdict'][0]}/{fixed['verdict'][1] if fixed['verdict'][0] == 'exc' else '-'}"
        )
        assert floor >= len(payload), f"G-A1 {key}: F={floor} < B={len(payload)}"
        assert floor <= 8 * len(payload) + 4 * MIB, f"G-A2 {key}: F={floor} 超过 8B+4MiB"
        assert fixed["peak"] == max(fixed["peaks"]), key

    # 历史分解说明（**非** `F(x)` 来源）：地板由 decode（1×）+ StringIO（4×）构成。
    payload = _DATA_CACHE["S-W6a-4000"]

    def decode_only():
        return payload.decode("utf-8-sig")

    decoded = decode_only()
    verdict_decode, peak_decode, _ = _measure_once(decode_only, trace=True)
    verdict_buffer, peak_buffer, _ = _measure_once(
        lambda: io.StringIO(decoded, newline=""), trace=True
    )
    verdict_loop, peak_loop, _ = _measure_once(
        lambda: sum(1 for _ in csv.reader(io.StringIO(decoded, newline=""))), trace=True
    )
    print(
        f"[floor] S-W6a-4000 bytes={len(payload)} decode_peak={peak_decode / MIB:.3f}MiB "
        f"({peak_decode / len(payload):.2f}x) StringIO_peak={peak_buffer / MIB:.3f}MiB "
        f"({peak_buffer / len(payload):.2f}x) reader_loop_peak={peak_loop / MIB:.3f}MiB "
        f"({peak_loop / len(payload):.2f}x) verdicts="
        f"{verdict_decode[0]}/{verdict_buffer[0]}/{verdict_loop[0]}"
    )


def test_td_f038_m4e_t1c_width_insensitivity():
    """TD-F038-M4（T-1c，§14.2，新增收紧）：列维度不敏感。

    `work(S-W6b) ≤ 1.5 × work(S-W6c) + 2 MiB`。S-W6b = `wide_csv(5001, 651,
    nonblank=True)`、S-W6c = `wide_csv(5001, 14, nonblank=True)`——两者行数相同（均
    超过 5000），仅宽度差 637 列；保留量不随宽度增长 ⇒ 两侧均必须为
    **413 `IMPORT_ROW_LIMIT_EXCEEDED{limit_rows:5000}`**（超额路径不保留内容）。
    """
    assert ImportRowLimitExceeded.status_code == 413
    assert ImportRowLimitExceeded.code == "IMPORT_ROW_LIMIT_EXCEEDED"
    expected_verdict = ("exc", "ImportRowLimitExceeded", {"limit_rows": CAP_ROWS})

    wide = _sample("S-W6b")
    narrow = _sample("S-W6c")
    assert wide["verdict"] == expected_verdict, f"S-W6b verdict={wide['verdict']}"
    assert narrow["verdict"] == expected_verdict, f"S-W6c verdict={narrow['verdict']}"

    work_wide = _work("S-W6b", wide)
    work_narrow = _work("S-W6c", narrow)
    bound = 1.5 * work_narrow + 2 * MIB
    print(
        f"[T-1c] S-W6b(651宽) B={wide['bytes']} peak={wide['peak'] / MIB:.3f}MiB "
        f"work={work_wide / MIB:.3f}MiB | S-W6c(14宽) B={narrow['bytes']} "
        f"peak={narrow['peak'] / MIB:.3f}MiB work={work_narrow / MIB:.3f}MiB "
        f"bound={bound / MIB:.3f}MiB verdicts={wide['verdict'][1]}/{narrow['verdict'][1]}"
    )
    assert work_wide <= bound, (
        f"T-1c: work(S-W6b)={work_wide / MIB:.3f}MiB > "
        f"1.5×work(S-W6c)+2MiB={bound / MIB:.3f}MiB"
    )


def test_td_f038_m4c_t6_report_only():
    """TD-F038-M4（T-6，§13.2 报告项，不判 FAIL）：S-W2 4N 峰值 ≤ 8 MiB。"""
    large = _sample("S-W2-4N10001")
    print(
        f"[T-6] S-W2 4N fixed_peak={large['peak'] / MIB:.3f}MiB "
        f"(report only, ≤8MiB, within={large['peak'] <= 8 * MIB})"
    )
    assert large["peak"] == max(large["peaks"])


@pytest.mark.parametrize("key", ["S-W1-1001", "S-W2-1001", "S-W6a-4000"])
def test_td_f038_m5_time_not_degraded(key):
    """TD-F038-M5（T-5 / ON-5）：记录秒数，`fixed_time ≤ baseline_time`（无绝对上限）。"""
    baseline, fixed = _pair(key)
    print(
        f"[T-5] {key}: t_fixed_traced={fixed['time_traced']:.3f}s "
        f"t_baseline_traced={baseline['time_traced']:.3f}s | "
        f"t_fixed_raw={fixed['time_raw']:.3f}s t_baseline_raw={baseline['time_raw']:.3f}s"
    )
    assert fixed["time_traced"] <= baseline["time_traced"]
    assert fixed["time_raw"] <= baseline["time_raw"]


def test_td_f038_m6_repeatability_and_ru_maxrss():
    """TD-F038-M6（A12 / R-11 / §14.3 G-A3）：两次重复值 + `B/F/F/B/peak/work` 双侧。

    断言仍用两次的**较大值**（保守）；`ru_maxrss` 仅报告（含 lxml C 层分配）。
    """
    import resource as resource_module

    results = {
        key: {"baseline": _pair(key)[0], "fixed": _pair(key)[1]}
        for key in ("S-W1-1001", "S-W2-1001", "S-W6a-4000", "legal-5000")
    }
    for key, sides in results.items():
        payload = _data(key)
        floor = floor_bytes(payload, _SAMPLES[key][0])
        for side, result in sides.items():
            assert len(result["peaks"]) == 2, key
            assert result["peak"] == max(result["peaks"])
            assert len(result["time_traced_runs"]) == 2
            work = excess(result, floor=floor)
            print(
                f"[M6] {key:12s} {side:8s} B={result['bytes']:9d} F={floor / MIB:9.3f}MiB "
                f"F/B={floor / result['bytes']:.3f} "
                f"peaks={[round(p / MIB, 3) for p in result['peaks']]} "
                f"peak={result['peak'] / MIB:9.3f}MiB work={work / MIB:9.3f}MiB"
            )
    peak_rss = resource_module.getrusage(resource_module.RUSAGE_SELF).ru_maxrss
    print(f"[M6] ru_maxrss={peak_rss}KiB (report only; includes lxml C allocations)")


# =============================================================================
# S 层（静态守卫）
# =============================================================================


def _git(*args: str) -> str:
    """只读 git 调用（`diff` / `status` / `rev-parse`）。"""
    completed = subprocess.run(
        ["git", "-C", str(REPO_ROOT), *args], capture_output=True, text=True, check=True
    )
    return completed.stdout


def test_td_f038_s1_zero_diff_and_frozen_contract():
    """TD-F038-S1（A13 / B12 / N2–N5）：契约/数据库/前端/依赖零 diff + 冻结常量。"""
    zero_diff_paths = (
        "docs/api",
        "docs/database",
        "backend/alembic",
        "frontend",
        "backend/pyproject.toml",
        "backend/Dockerfile",
        "backend/tests/integration/test_f037_import_parse_errors.py",
    )
    diff = _git("diff", "--stat", START_COMMIT, "--", *zero_diff_paths)
    assert diff.strip() == "", f"禁止写入的路径出现 diff:\n{diff}"

    templates = {item["resource_type"]: item for item in resource_import.templates_payload()}
    assert [column["name"] for column in templates["BARE_METAL"]["columns"]] == BM_HEADERS
    assert len(templates["VIRTUAL_MACHINE"]["columns"]) == 10
    assert len(unified.COLUMN_NAMES) == 26
    assert list(unified.COLUMN_NAMES) == W26_HEADERS

    settings = Settings(_env_file=None)
    assert (
        settings.import_max_file_bytes,
        settings.import_max_rows,
        settings.import_error_preview_limit,
    ) == (CAP_BYTES, CAP_ROWS, CAP_PREVIEW)


def test_td_f038_s2_production_call_sites_pass_both_keywords():
    """TD-F038-S2 / E-1（R-9）：`parse_records(` 恰 2 处且均带 `max_rows` / `keep_cells`。"""
    call_sites: list[tuple[str, int, list[str], str]] = []
    for path in sorted((BACKEND_DIR / "app").rglob("*.py")):
        tree = ast.parse(path.read_text())
        for node in ast.walk(tree):
            if (
                isinstance(node, ast.Call)
                and isinstance(node.func, ast.Name)
                and node.func.id == "parse_records"
            ):
                call_sites.append(
                    (
                        str(path.relative_to(BACKEND_DIR)),
                        node.lineno,
                        sorted(keyword.arg for keyword in node.keywords),
                        ast.unparse(node),
                    )
                )
            elif (
                isinstance(node, ast.Call)
                and isinstance(node.func, ast.Attribute)
                and node.func.attr == "parse_records"
            ):
                call_sites.append(
                    (
                        str(path.relative_to(BACKEND_DIR)),
                        node.lineno,
                        sorted(keyword.arg for keyword in node.keywords),
                        ast.unparse(node),
                    )
                )

    assert len(call_sites) == 2, call_sites
    for path, _lineno, keywords, _call in call_sites:
        assert keywords == ["keep_cells", "max_rows"], (path, keywords)
    joined = " ".join(item[3] for item in call_sites)
    assert "_MAX_TEMPLATE_COLUMNS" in joined
    assert "len(COLUMN_NAMES)" in joined
    assert resource_import._MAX_TEMPLATE_COLUMNS == max(
        len(template.columns) for template in resource_import.TEMPLATES
    )


def test_td_f038_s3_exception_surface_unchanged():
    """TD-F038-S3（R-8 / R-10）：无新增 `except`、白名单 8 类、无严格 CSV / 新依赖。"""
    assert len(resource_import._PARSE_FAILURES) == 8
    source_path = BACKEND_DIR / "app/services/resource_import.py"
    unified_path = BACKEND_DIR / "app/services/resource_interface_ip_import.py"
    tree = ast.parse(source_path.read_text())
    handlers = [node for node in ast.walk(tree) if isinstance(node, ast.ExceptHandler)]
    broad = [
        node
        for node in handlers
        if isinstance(node.type, ast.Name) and node.type.id == "Exception"
    ]
    assert len(broad) == 2, "`except Exception` 必须仍为 2 处（_close_workbook / write_import）"

    baseline_handlers = None
    if Path(BASELINE_PATH).exists():
        baseline_tree = ast.parse(Path(BASELINE_PATH).read_text())
        baseline_handlers = sum(
            1
            for node in ast.walk(baseline_tree)
            if isinstance(node, ast.ExceptHandler)
        )
        assert len(handlers) == baseline_handlers, "本次 diff 不得新增 except 子句"

    text_source = source_path.read_text()
    text_unified = unified_path.read_text()
    for banned in ("field_size_limit", "strict=True", "csv.QUOTE"):
        assert banned not in text_source, banned
        assert banned not in text_unified, banned
    assert "lxml" not in text_unified

    pyproject = (BACKEND_DIR / "pyproject.toml").read_text()
    assert "lxml" not in pyproject, "lxml 必须保持 openpyxl 的可选加速依赖（不得声明）"

    # `OSError` 只允许出现在 `_parse_xlsx` 行区间内
    functions = {
        node.name: (node.lineno, node.end_lineno)
        for node in ast.walk(tree)
        if isinstance(node, ast.FunctionDef)
    }
    start, end = functions["_parse_xlsx"]
    lines = text_source.splitlines()
    for index, line in enumerate(lines, start=1):
        if "OSError" in line:
            assert start <= index <= end, f"OSError 出现在 _parse_xlsx 之外：{index}: {line}"

    # `_invalid_file_on_parse_failure` 子句保持不变
    guard = next(
        node
        for node in ast.walk(tree)
        if isinstance(node, ast.FunctionDef) and node.name == "_invalid_file_on_parse_failure"
    )
    assert "AppError" in ast.unparse(guard)
    assert "failures" in ast.unparse(guard)


def test_td_f038_s4_no_full_materialisation_left():
    """TD-F038-S4（R-7）：无 `list(csv.reader(` / `list(worksheet.iter_rows(` 残留。"""
    for name in ("resource_import.py", "resource_interface_ip_import.py"):
        text_source = (BACKEND_DIR / "app/services" / name).read_text()
        for banned in ("list(csv.reader(", "list(worksheet.iter_rows("):
            assert banned not in text_source, f"{name}: {banned}"

    tree = ast.parse((BACKEND_DIR / "app/services/resource_import.py").read_text())
    collect = next(
        node
        for node in ast.walk(tree)
        if isinstance(node, ast.FunctionDef) and node.name == "_collect_records"
    )
    loop = next(node for node in ast.walk(collect) if isinstance(node, ast.For))
    exits = [
        node
        for node in ast.walk(loop)
        if isinstance(node, (ast.Break, ast.Return, ast.Continue))
        and isinstance(node, (ast.Break, ast.Return))
    ]
    assert exits == [], "`_collect_records` 的 for 体内不得 break / return"
    assert ast.unparse(loop).count("continue") >= 1

    # 禁止整行 list() 复制（IMP-2）
    assert "list(record)" not in ast.unparse(collect)


def test_td_f038_s5_l7_double_prime_no_io_no_except():
    """TD-F038-S5（A14 / L7″）：新增自有代码内无 try/except/I/O/日志。"""
    tree = ast.parse((BACKEND_DIR / "app/services/resource_import.py").read_text())
    targets = ("_is_blank_record", "_collect_records", "_finalise", "parse_records")
    for node in ast.walk(tree):
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name in targets:
            body = ast.unparse(node)
            for banned in ("try:", "except", "open(", "Path(", "NamedTemporaryFile", "logger."):
                assert banned not in body, f"{node.name}: {banned}"
            names = {
                item.id
                for item in ast.walk(node)
                if isinstance(item, ast.Name)
            }
            assert "logger" not in names

    # `_cell_text` 调用只允许出现在 `_finalise`
    for node in ast.walk(tree):
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name in targets:
            calls = [
                item
                for item in ast.walk(node)
                if isinstance(item, ast.Call)
                and isinstance(item.func, ast.Name)
                and item.func.id == "_cell_text"
            ]
            if node.name == "_finalise":
                assert calls == [
                    item
                    for item in ast.walk(node)
                    if isinstance(item, ast.Call)
                    and isinstance(item.func, ast.Name)
                    and item.func.id == "_cell_text"
                ]
            else:
                assert calls == [], f"{node.name} 不得构造文本"


def test_td_f038_s6_change_surface_is_two_services():
    """TD-F038-S6（§13.6 / §13.10 放行写入）：仅两个 service 有实现 diff。"""
    changed = [
        line.strip()
        for line in _git("diff", "--name-only", START_COMMIT, "--", "backend/app")
        .strip()
        .splitlines()
        if line.strip()
    ]
    assert sorted(changed) == [
        "backend/app/services/resource_import.py",
        "backend/app/services/resource_interface_ip_import.py",
    ], changed

    status = _git("status", "--porcelain", "--untracked-files=all")
    allowed = {
        "backend/app/services/resource_import.py",
        "backend/app/services/resource_interface_ip_import.py",
        "backend/tests/integration/test_f038_parse_bounds.py",
    }
    for line in status.splitlines():
        match = re.match(r"^(..) (.+)$", line)
        assert match is not None, line
        path = match.group(2).strip()
        assert path in allowed or path.startswith("docs/features/F038/"), line
